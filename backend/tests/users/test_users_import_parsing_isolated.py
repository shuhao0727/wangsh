"""Pure import-file parsing regressions; run only with the offline bootstrap."""

import asyncio
import io
import zipfile
from contextlib import asynccontextmanager

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook
from sqlalchemy import event, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.api.endpoints.management.users import import_service
from app.models import User


def _xlsx(headers=None, rows=None):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(headers or ["学号", "姓名"])
    for row in rows if rows is not None else [["B1-001", "测试学生"]]:
        worksheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _broken_worksheet():
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_xlsx())) as source, zipfile.ZipFile(output, "w") as target:
        for member in source.infolist():
            content = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                content = b"<worksheet><malformed>"
            target.writestr(member, content)
    return output.getvalue()


class NoWritesDb:
    async def rollback(self):
        pass

    def __getattr__(self, name):
        raise AssertionError(f"Invalid file must not reach DB operation: {name}")


@pytest.mark.parametrize(
    ("filename", "content"),
    [
        ("users.csv", b"\xff\xfe\xff"),
        ("users.txt", b"\xff\xfe\xff"),
        ("users.xlsx", b"corrupt"),
        ("users.xlsx", _xlsx()[:100]),
        ("users.xlsx", _broken_worksheet()),
    ],
)
def test_invalid_import_file_is_client_error_before_db_write(filename, content):
    file = UploadFile(filename=filename, file=io.BytesIO(content))
    with pytest.raises(HTTPException) as caught:
        asyncio.run(import_service.import_users(file, {"id": 1, "role_code": "admin"}, NoWritesDb()))
    assert caught.value.status_code == 400
    assert caught.value.detail


def test_zip_missing_xlsx_members_is_client_error():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("README.txt", "This ZIP is not an Excel workbook")
    with pytest.raises(HTTPException) as caught:
        import_service.parse_import_rows("users.xlsx", output.getvalue())
    assert caught.value.status_code == 400


def test_declared_workbook_missing_from_archive_is_client_error():
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_xlsx())) as source, zipfile.ZipFile(output, "w") as target:
        for member in source.infolist():
            if member.filename != "xl/workbook.xml":
                target.writestr(member, source.read(member.filename))
    with pytest.raises(HTTPException) as caught:
        import_service.parse_import_rows("users.xlsx", output.getvalue())
    assert caught.value.status_code == 400


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_valid_import_rows_preserve_text_student_ids(format):
    content = "学号,姓名\n00017,测试学生\n".encode("utf-8-sig") if format == "csv" else _xlsx()
    headers, rows = import_service.parse_import_rows(f"users.{format}", content)
    assert headers == ["学号", "姓名"]
    row = next(rows)
    assert row == {"学号": "00017" if format == "csv" else "B1-001", "姓名": "测试学生"}


def test_missing_required_header_keeps_existing_client_error():
    with pytest.raises(HTTPException) as caught:
        import_service._validate_import_headers(["姓名"])
    assert caught.value.status_code == 400


def test_unexpected_parser_failure_is_not_masked_as_bad_input(monkeypatch):
    def unavailable(*args, **kwargs):
        raise RuntimeError("unexpected parser infrastructure failure")

    monkeypatch.setattr(import_service, "load_workbook", unavailable)
    with pytest.raises(RuntimeError, match="unexpected parser"):
        import_service.parse_import_rows("users.xlsx", _xlsx())


def _lazy_broken(at_header, content=None):
    with zipfile.ZipFile(io.BytesIO(content or _xlsx())) as source:
        members = {name: source.read(name) for name in source.namelist()}
    name = "xl/worksheets/sheet1.xml"
    marker = b"<sheetData>" if at_header else b"</sheetData>"
    replacement = b" " * 32768 + (b"<sheetData><invalid></wrong>" if at_header else b"<invalid></wrong></sheetData>")
    members[name] = members[name].replace(marker, replacement)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name, content in members.items():
            target.writestr(name, content)
    return output.getvalue()


def _track_workbooks(monkeypatch):
    tracked = []
    actual = import_service.load_workbook

    def load(*args, **kwargs):
        workbook = actual(*args, **kwargs)
        tracked.append((workbook, kwargs["filename"]))
        return workbook

    monkeypatch.setattr(import_service, "load_workbook", load)
    return tracked


def _assert_closed(tracked):
    assert len(tracked) == 1
    workbook, stream = tracked[0]
    assert workbook._archive.fp is None
    assert stream.closed


def test_lazy_header_parse_error_is_400_and_closes_loaded_workbook(monkeypatch):
    tracked = _track_workbooks(monkeypatch)
    with pytest.raises(HTTPException) as caught:
        import_service.parse_import_rows("users.xlsx", _lazy_broken(at_header=True))
    assert caught.value.status_code == 400
    _assert_closed(tracked)


def test_lazy_tail_parse_error_rolls_back_and_closes_loaded_workbook(monkeypatch):
    tracked = _track_workbooks(monkeypatch)
    observed = []

    async def consume(db, reader, current_user):
        for row in reader:
            observed.append(row)
        raise AssertionError("Malformed tail must fail before commit")

    class Db:
        rollback_count = 0

        async def rollback(self):
            self.rollback_count += 1

        async def commit(self):
            raise AssertionError("Invalid workbook must not commit")

    db = Db()
    monkeypatch.setattr(import_service, "_process_import_rows", consume)
    file = UploadFile(filename="users.xlsx", file=io.BytesIO(_lazy_broken(at_header=False)))
    with pytest.raises(HTTPException) as caught:
        asyncio.run(import_service.import_users(file, {"id": 1}, db))
    assert caught.value.status_code == 400
    assert len(observed) == 1
    assert db.rollback_count == 1
    _assert_closed(tracked)


def test_normal_exhaustion_closes_loaded_workbook(monkeypatch):
    tracked = _track_workbooks(monkeypatch)
    _, rows = import_service.parse_import_rows("users.xlsx", _xlsx())
    assert len(list(rows)) == 1
    _assert_closed(tracked)


def test_required_header_rejection_closes_loaded_workbook(monkeypatch):
    tracked = _track_workbooks(monkeypatch)
    file = UploadFile(filename="users.xlsx", file=io.BytesIO(_xlsx(["姓名"])))
    with pytest.raises(HTTPException) as caught:
        asyncio.run(import_service.import_users(file, {"id": 1}, NoWritesDb()))
    assert caught.value.status_code == 400
    _assert_closed(tracked)


def test_early_processing_failure_keeps_500_and_closes_loaded_workbook(monkeypatch):
    tracked = _track_workbooks(monkeypatch)

    async def fail(db, reader, current_user):
        raise RuntimeError("unexpected database/infrastructure failure")

    monkeypatch.setattr(import_service, "_process_import_rows", fail)
    file = UploadFile(filename="users.xlsx", file=io.BytesIO(_xlsx()))
    with pytest.raises(HTTPException) as caught:
        asyncio.run(import_service.import_users(file, {"id": 1}, NoWritesDb()))
    assert caught.value.status_code == 500
    _assert_closed(tracked)


def test_cancellation_is_not_converted_to_400_and_closes_loaded_workbook(monkeypatch):
    tracked = _track_workbooks(monkeypatch)

    async def cancel(db, reader, current_user):
        raise asyncio.CancelledError()

    monkeypatch.setattr(import_service, "_process_import_rows", cancel)
    file = UploadFile(filename="users.xlsx", file=io.BytesIO(_xlsx()))
    with pytest.raises(asyncio.CancelledError):
        asyncio.run(import_service.import_users(file, {"id": 1}, NoWritesDb()))
    _assert_closed(tracked)


def test_programming_keyerror_is_not_converted_to_400(monkeypatch):
    def bug(*args, **kwargs):
        raise KeyError("implementation bug unrelated to ZIP member lookup")

    monkeypatch.setattr(import_service, "load_workbook", bug)
    with pytest.raises(KeyError, match="implementation bug"):
        import_service.parse_import_rows("users.xlsx", _xlsx())


@asynccontextmanager
async def _import_sessions():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")

    @event.listens_for(engine.sync_engine, "connect")
    def disable_legacy_transactions(connection, _record):
        connection.isolation_level = None

    @event.listens_for(engine.sync_engine, "begin")
    def begin_transaction(connection):
        # SAVEPOINT RELEASE must not implicitly commit the SQLite outer txn.
        connection.exec_driver_sql("BEGIN")

    sessions = async_sessionmaker(engine, expire_on_commit=False, autoflush=False)
    try:
        async with engine.begin() as connection:
            await connection.run_sync(User.__table__.create)
        async with sessions() as db:
            db.add_all([
                User(id=1, username="import_admin", full_name="合成管理员", role_code="admin"),
                User(id=2, student_id="B1-001", full_name="原姓名", role_code="student"),
                User(id=3, username="retired_name", full_name="已删除", is_deleted=True),
            ])
            await db.commit()
        yield sessions
    finally:
        await engine.dispose()


async def _saved_users(sessions):
    async with sessions() as db:
        return (await db.execute(
            select(User.student_id, User.full_name).order_by(User.id)
        )).all()


def test_lazy_tail_error_rolls_back_real_inserts_and_updates(monkeypatch):
    tracked = _track_workbooks(monkeypatch)
    processed = []
    actual = import_service._import_user_row

    async def track_row(*args, **kwargs):
        result = await actual(*args, **kwargs)
        processed.append(result)
        return result

    monkeypatch.setattr(import_service, "_import_user_row", track_row)

    async def run():
        async with _import_sessions() as sessions:
            before = await _saved_users(sessions)
            content = _xlsx(rows=[["B1-001", "不应保存的更新"], ["B1-002", "不应保存的新建"]])
            upload = UploadFile(filename="users.xlsx", file=io.BytesIO(_lazy_broken(False, content)))
            async with sessions() as db:
                with pytest.raises(HTTPException) as caught:
                    await import_service.import_users(upload, {"id": 1}, db)
                assert caught.value.status_code == 400
                assert not db.in_transaction()
            assert await _saved_users(sessions) == before

    asyncio.run(run())
    assert [created for _result, created in processed] == [False, True]
    _assert_closed(tracked)


def test_cancelled_import_rolls_back_before_session_is_closed(monkeypatch):
    tracked = _track_workbooks(monkeypatch)
    actual = import_service._import_user_row

    async def cancel_third_row(db, row, current_user, row_number):
        if row_number == 3:
            asyncio.current_task().cancel()
            await asyncio.sleep(0)
        return await actual(db, row, current_user, row_number)

    monkeypatch.setattr(import_service, "_import_user_row", cancel_third_row)

    async def run():
        async with _import_sessions() as sessions:
            before = await _saved_users(sessions)
            upload = UploadFile(filename="users.xlsx", file=io.BytesIO(
                _xlsx(rows=[["B1-001", "取消的更新"], ["B1-002", "取消的新建"], ["B1-003", "触发取消"]])
            ))
            async with sessions() as db:
                task = asyncio.create_task(import_service.import_users(upload, {"id": 1}, db))
                with pytest.raises(asyncio.CancelledError):
                    await task
                # Deliberately keep the session open: closing it would hide
                # missing rollback in the service behind dependency cleanup.
                assert not db.in_transaction()
                await db.commit()
            assert await _saved_users(sessions) == before

    asyncio.run(run())
    _assert_closed(tracked)


def test_final_commit_failure_rolls_back_all_rows_and_closes_workbook(monkeypatch):
    tracked = _track_workbooks(monkeypatch)

    async def fail_commit():
        raise RuntimeError("synthetic final commit failure")

    async def run():
        async with _import_sessions() as sessions:
            before = await _saved_users(sessions)
            upload = UploadFile(filename="users.xlsx", file=io.BytesIO(
                _xlsx(rows=[["B1-001", "失败的更新"], ["B1-002", "失败的新建"]])
            ))
            async with sessions() as db:
                monkeypatch.setattr(db, "commit", fail_commit)
                with pytest.raises(HTTPException) as caught:
                    await import_service.import_users(upload, {"id": 1}, db)
                assert caught.value.status_code == 500
                assert not db.in_transaction()
            assert await _saved_users(sessions) == before

    asyncio.run(run())
    _assert_closed(tracked)


def test_each_successful_row_releases_its_savepoint():
    async def run():
        async with _import_sessions() as sessions:
            async with sessions() as db:
                rows = iter({"学号": f"B1-{i:03}", "姓名": f"学生{i}"} for i in range(2, 10))
                results, imported, updated = await import_service._process_import_rows(db, rows, {"id": 1})
                assert len(results) == imported == 8
                assert updated == 0
                assert not db.in_nested_transaction()
                assert db.in_transaction()
                await db.rollback()
            assert len(await _saved_users(sessions)) == 3

    asyncio.run(run())


def test_row_validation_and_sql_errors_keep_other_successful_rows(monkeypatch):
    tracked = _track_workbooks(monkeypatch)

    async def run():
        async with _import_sessions() as sessions:
            upload = UploadFile(filename="users.xlsx", file=io.BytesIO(_xlsx(
                headers=["学号", "姓名", "用户名"],
                rows=[
                    ["B1-001", "更新成功", ""],
                    ["B1-001", "失败行不可覆盖", "import_admin"],
                    ["B1-002", "SQL唯一约束失败", "retired_name"],
                    ["B1-003", "", ""],
                    ["B1-004", "继续成功", "new_import_name"],
                ],
            )))
            async with sessions() as db:
                result = await import_service.import_users(upload, {"id": 1}, db)
                assert not db.in_transaction()
            assert not result.success
            assert (result.total_rows, result.imported_count, result.updated_count, result.error_count) == (5, 1, 1, 3)
            assert [error.row_number for error in result.errors] == [2, 3, 4]
            assert await _saved_users(sessions) == [
                (None, "合成管理员"), ("B1-001", "更新成功"), (None, "已删除"), ("B1-004", "继续成功"),
            ]

    asyncio.run(run())
    _assert_closed(tracked)
