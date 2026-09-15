"""XBK XLSX text contract: real ZIP/workbook round trips, synthetic query rows only.

No formula evaluator, Excel application, production DB or business workbook is used.
Run in an isolated app-import harness (see backend/tests/README.md); CaptureDb
only supplies query results, while the actual endpoint/builders/openpyxl run.
"""
import asyncio
from io import BytesIO
from types import SimpleNamespace
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils.protection import hash_password

from app.api.endpoints.xbk import exports, import_export
from app.core.config import settings
from app.services.xbk.exports import common
from app.services.xbk.course_selection_workbook import METADATA_SHEET


NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
YEAR = "2026-2027"
TERM = "上学期"
# Harmless arithmetic strings only. No external links, DDE, commands or macros.
PAYLOADS = [
    prefix + operator + "1+1" + suffix
    for prefix, suffix in [("", ""), (" ", " "), ("\t", "\t"), ("\n", "\n"),
                           ("\r\n", "\r\n"), ("\u00a0", "\u2003"), ("\ufeff", "\u200b")]
    for operator in "=+-@"
] + ["=", "+", "-", "@", " ", "\t", "\n", "\r", "\r\n", "", "0012", "0007",
     "00000000000000000000000000000000000000000000000001", "1.50", "1E10", "-12",
     "中文é🙂全角＝＋－＠", "组合e\u0301", "含\t制表\n换行\r回车", "正常文本",
     "#N/A", "#REF!", "#DIV/0!", "'0012"]
SCOPES = ["students", "courses", "selections", "course_results", "unselected", "suspended"]
KINDS = ["course-selection", "distribution", "teacher-distribution"]
STUDENT_FIELDS = [("学年", "year"), ("学期", "term"), ("年级", "grade"),
                  ("班级", "class_name"), ("学号", "student_no"), ("姓名", "name"), ("性别", "gender")]
FIELDS = {
    "students": STUDENT_FIELDS, "unselected": STUDENT_FIELDS, "suspended": STUDENT_FIELDS,
    "courses": [("学年", "year"), ("学期", "term"), ("年级", "grade"), ("课程代码", "course_code"),
                ("课程名称", "course_name"), ("课程负责人", "teacher"), ("各班限报人数", "quota"), ("上课地点", "location")],
    "selections": [("学年", "year"), ("学期", "term"), ("年级", "grade"), ("学号", "student_no"),
                   ("姓名", "student_name"), ("课程代码", "course_code")],
    "course_results": [("学年", "year"), ("学期", "term"), ("年级", "grade"), ("班级", "class_name"),
                       ("学号", "student_no"), ("姓名", "student_name"), ("课程代码", "course_code"),
                       ("课程名称", "course_name"), ("负责人", "teacher"), ("地点", "location")],
}


class Rows:
    def __init__(self, rows):
        self.rows = rows

    def scalars(self):
        return self

    def all(self):
        return self.rows


class CaptureDb:
    """No engine/session/config access, and no writes are available."""
    def __init__(self, *rows):
        self.results = iter(rows)
        self.calls = []

    async def execute(self, stmt):
        self.calls.append(str(stmt))
        return Rows(next(self.results))


def synthetic_row(**changes):
    row = dict(year=YEAR, term=TERM, grade="高一", class_name="1班", student_no="0012",
               name="合成学生甲", gender="女", course_code="0007", course_name="合成课程乙",
               teacher="合成教师丙", location="合成教室丁", student_name="合成姓名快照戊", quota=3)
    row.update(changes)
    return SimpleNamespace(**row)


async def response_bytes(response):
    return b"".join([chunk async for chunk in response.body_iterator])


def ordinary_bytes(scope, row, file_format="xlsx", *, diagnostic_rows=()):
    results = ([] if row is None else [row],)
    if scope in {"selections", "course_results"}:
        results += (list(diagnostic_rows),)
    async def run():
        response = await import_export.export_data(
            scope=scope, year=YEAR, term=TERM, grade=None, class_name=None, search_text=None,
            format=file_format, db=CaptureDb(*results), _={"role": "admin"})
        assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        return await response_bytes(response)
    return asyncio.run(run())


def specialized_bytes(kind, row, *, grade=None, term=TERM, empty=False):
    results = ([] if empty else [row],)
    if kind != "distribution":
        results += ([] if empty else [row],)
    if kind == "course-selection" and not empty:
        # course_selection now populates the class sheet from active selections.
        results += ([row],)

    async def run():
        response = await exports.export_tables(
            export_type=kind, year=YEAR, term=term, grade=grade, class_name=None,
            year_start=None, year_end=None, db=CaptureDb(*results), _={"role": "admin"})
        assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        return await response_bytes(response)
    return asyncio.run(run())


def assert_xlsx_contract(content):
    """Independently check serialized cell types/no <f>, and reader-visible types.

    Template validation formula1 is trusted application logic, not a cell <f>.
    Its preservation is separately asserted below; no formula is ever executed.
    """
    with ZipFile(BytesIO(content)) as archive:
        assert not any("externalLink" in path or "vbaProject" in path for path in archive.namelist())
        for path in archive.namelist():
            if not (path.startswith("xl/worksheets/sheet") and path.endswith(".xml")):
                continue
            tree = ET.fromstring(archive.read(path))
            assert not tree.findall(".//s:c/s:f", NS), path
            for cell in tree.findall(".//s:c", NS):
                assert cell.get("t") not in {"f", "e", "str"}, (path, cell.attrib)
    wb = load_workbook(BytesIO(content), data_only=False)
    for ws in wb:
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    assert cell.data_type == "s", (ws.title, cell.coordinate, cell.value)
                    assert cell.number_format == "@", (ws.title, cell.coordinate, cell.number_format)
    return wb


def assert_text(cell, expected):
    if expected in ("", None):
        # openpyxl reloads empty inline strings as None, not an invented apostrophe.
        assert cell.value is None
        assert cell.data_type != "f"
    else:
        assert cell.value == expected, (cell.coordinate, cell.value, expected)
        assert cell.data_type == "s"
        assert cell.number_format == "@"


@pytest.mark.parametrize("scope", SCOPES)
@pytest.mark.parametrize("payload", PAYLOADS)
def test_ordinary_export_every_text_field(scope, payload):
    text_fields = [attr for _, attr in FIELDS[scope] if attr != "quota"]
    row = synthetic_row(**dict.fromkeys(text_fields, payload))
    ws = assert_xlsx_contract(ordinary_bytes(scope, row))["data"]
    assert [cell.value for cell in ws[1]] == [header for header, _ in FIELDS[scope]]
    for col, (_, attr) in enumerate(FIELDS[scope], start=1):
        cell = ws.cell(2, col)
        if attr == "quota":
            assert (cell.value, cell.data_type) == (3, "n")
        else:
            assert_text(cell, payload)


@pytest.mark.parametrize("scope", SCOPES)
def test_ordinary_export_distinct_field_mapping_and_legacy_xls_alias(scope):
    row = synthetic_row()
    ws = assert_xlsx_contract(ordinary_bytes(scope, row, file_format="xls"))["data"]
    for col, (_, attr) in enumerate(FIELDS[scope], start=1):
        if attr != "quota":
            assert_text(ws.cell(2, col), getattr(row, attr))
        else:
            assert (ws.cell(2, col).value, ws.cell(2, col).data_type) == (3, "n")


@pytest.mark.parametrize("kind", KINDS)
@pytest.mark.parametrize("payload", PAYLOADS)
def test_specialized_export_every_body_text_field(kind, payload):
    fields = ["course_code", "course_name", "teacher", "location", "name", "student_name", "student_no"]
    row = synthetic_row(**dict.fromkeys(fields, payload))
    wb = assert_xlsx_contract(specialized_bytes(kind, row, grade="高一"))
    if kind == "course-selection":
        for coord in ["A3", "B3", "C3", "E3"]:
            assert_text(wb.worksheets[0][coord], payload)
        assert (wb.worksheets[0]["D3"].value, wb.worksheets[0]["D3"].data_type) == (3, "n")
        for coord in ["B2", "C2"]:
            assert_text(wb.worksheets[1][coord], payload)
        assert_text(wb.worksheets[1]["A2"], "1班")
    elif kind == "distribution":
        for cell in wb.active[3]:
            assert_text(cell, payload)
    else:
        assert_text(wb.active["B6"], payload)
        assert_text(wb.active["A2"], f"课程代码: {payload}  课程名称: {payload}  "
                    f"课程负责人: {payload}  上课地点: {payload}  人数: 1")


@pytest.mark.parametrize("kind", KINDS)
@pytest.mark.parametrize("payload", ["=1+1", "+1+1", "-1+1", "@1+1", " \t=1+1\r\n", "0012", "中文🙂", "#N/A"])
@pytest.mark.parametrize("filtered", [True, False])
def test_specialized_export_grade_class_and_title_text(kind, payload, filtered):
    row = synthetic_row(grade=payload, class_name=payload)
    wb = assert_xlsx_contract(specialized_bytes(kind, row, grade=payload if filtered else None, term=payload))
    if kind == "course-selection":
        assert_text(wb.worksheets[1]["A2"], payload)
        assert_text(wb.worksheets[1]["B2"], "0012")
        assert_text(wb.worksheets[1]["C2"], row.name)
    elif kind == "teacher-distribution":
        assert_text(wb.active["A6"], payload if filtered else payload + payload)
        assert payload in wb.active["A1"].value
    else:
        # Class title has intentional pre-existing presentation formatting.
        assert payload in wb.active["A1"].value  # term is preserved verbatim


@pytest.mark.parametrize("kind", KINDS)
def test_specialized_export_empty_and_distinct_mapping(kind):
    assert_xlsx_contract(specialized_bytes(kind, synthetic_row(), empty=True))
    row = synthetic_row()
    wb = assert_xlsx_contract(specialized_bytes(kind, row))
    if kind == "course-selection":
        for coord, attr in [("A3", "course_code"), ("B3", "course_name"), ("C3", "teacher"), ("E3", "location")]:
            assert_text(wb.worksheets[0][coord], getattr(row, attr))
    elif kind == "distribution":
        for col, attr in enumerate(["course_code", "course_name", "teacher", "location", "student_name"], 1):
            assert_text(wb.active.cell(3, col), getattr(row, attr))
    else:
        assert_text(wb.active["A6"], row.grade + row.class_name)
        assert_text(wb.active["B6"], row.student_name)


@pytest.mark.parametrize("value", [0, 1, 3, 2147483647])
def test_numeric_quota_remains_numeric(value):
    row = synthetic_row(quota=value)
    for content, sheet, coord in [(ordinary_bytes("courses", row), "data", "G2"),
                                  (specialized_bytes("course-selection", row), "校本课程目录", "D3")]:
        wb = assert_xlsx_contract(content)
        assert (wb[sheet][coord].value, wb[sheet][coord].data_type) == (value, "n")
        with ZipFile(BytesIO(content)) as archive:
            cells = ET.fromstring(archive.read("xl/worksheets/sheet1.xml")).findall(".//s:c", NS)
            cell = next(cell for cell in cells if cell.get("r") == coord)
            assert cell.get("t") == "n"
            assert cell.find("s:v", NS).text == str(value)


@pytest.mark.parametrize("course_code", ["=1+1", "=1+1\r\n"])
def test_template_validation_and_protection_are_preserved(course_code):
    content = specialized_bytes("course-selection", synthetic_row(course_code=course_code))
    wb = assert_xlsx_contract(content)
    catalog = wb["校本课程目录"]
    roster = next(ws for ws in wb.worksheets if ws.title not in {"校本课程目录", METADATA_SHEET})
    assert wb[METADATA_SHEET].sheet_state == "veryHidden"
    assert catalog.protection.sheet and roster.protection.sheet
    expected_password_hash = hash_password(settings.XBK_EXPORT_SHEET_PASSWORD)
    assert catalog.protection.password == expected_password_hash
    assert roster.protection.password == expected_password_hash
    assert roster["B2"].protection.locked and not roster["D2"].protection.locked
    assert roster["D2"].value == course_code
    validation = roster.data_validations.dataValidation[0]
    assert validation.formula1 == (
        '=IFERROR(OR(D2="",D2="未选",AND(D2<>"",'
        'COUNTIF(D:D,D2)<=VLOOKUP(D2&"",\'校本课程目录\'!$A$3:$D$3,4,0))),FALSE)'
    )
    with ZipFile(BytesIO(content)) as archive:
        catalog_tree = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        roster_tree = ET.fromstring(archive.read("xl/worksheets/sheet2.xml"))
        assert catalog_tree.find(".//s:sheetProtection", NS).get("password") == expected_password_hash
        assert roster_tree.find(".//s:sheetProtection", NS).get("password") == expected_password_hash
        assert roster_tree.find(".//s:dataValidation/s:formula1", NS).text == validation.formula1


@pytest.mark.parametrize("payload", PAYLOADS)
def test_force_text_cells_xml_literal_preservation_and_idempotence(payload):
    wb = Workbook()
    ws = wb.active
    ws.append([payload, 0, 12, -12, 1.5, None, True])
    ws.merge_cells("A3:C3")
    ws["A3"] = payload
    ws["C1"].number_format = "0.00"
    common.force_text_cells(ws)
    common.force_text_cells(ws)
    stream = BytesIO()
    wb.save(stream)
    loaded = assert_xlsx_contract(stream.getvalue()).active
    assert_text(loaded["A1"], payload)
    assert_text(loaded["A3"], payload)
    for coordinate, value in [("B1", 0), ("C1", 12), ("D1", -12), ("E1", 1.5)]:
        assert (loaded[coordinate].value, loaded[coordinate].data_type) == (value, "n")
    assert loaded["C1"].number_format == "0.00"
    assert loaded["F1"].value is None
    assert (loaded["G1"].value, loaded["G1"].data_type) == (True, "b")
    with ZipFile(BytesIO(stream.getvalue())) as archive:
        tree = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        for coordinate in ["A1", "A3"]:
            cell = next(cell for cell in tree.findall(".//s:c", NS) if cell.get("r") == coordinate)
            assert cell.get("t") == "inlineStr"
            assert "".join(t.text or "" for t in cell.findall(".//s:t", NS)) == payload


@pytest.mark.parametrize("payload", [
    "\r", "\r\n", "a\rb\r\nc\nd\t", "&#13;\r_x000D_", " =1+1\r\n", "\r\r\n\r",
])
@pytest.mark.parametrize("destination", ["stream", "path"])
def test_cr_round_trip_repeated_save_is_local_to_export_workbook(payload, destination, tmp_path):
    # Run this suite in separate OPENPYXL_LXML=True/False processes. Do not
    # switch writer globals inside a process: openpyxl binds them at import.
    from openpyxl.cell import _writer
    from openpyxl.xml import functions

    original_save = Workbook.save
    original_write_cell = _writer.write_cell
    original_tostring = functions.tostring
    unrelated = Workbook()
    unrelated.active["A1"] = "=1+1"  # Never evaluated, even in the control workbook.
    wb = Workbook()
    ws = wb.active
    ws.append([payload, "0012", -12, 1.5, True, "&#13;_x000D_"])
    common.force_text_cells(ws)
    save_once = wb.save
    common.force_text_cells(ws)
    assert wb.save == save_once  # Finalizing twice must not stack save adapters.
    assert ws["A1"].value == payload
    assert Workbook.save is original_save
    assert _writer.write_cell is original_write_cell
    assert functions.tostring is original_tostring
    assert unrelated.save.__func__ is original_save
    assert unrelated.active["A1"].data_type == "f"

    for index in range(2):
        target = BytesIO() if destination == "stream" else tmp_path / f"synthetic-{index}.xlsx"
        wb.save(target)
        content = target.getvalue() if destination == "stream" else target.read_bytes()
        loaded = assert_xlsx_contract(content).active
        assert_text(loaded["A1"], payload)
        assert_text(loaded["B1"], "0012")
        assert_text(loaded["F1"], "&#13;_x000D_")
        for coord, value, data_type in [("C1", -12, "n"), ("D1", 1.5, "n"), ("E1", True, "b")]:
            assert (loaded[coord].value, loaded[coord].data_type) == (value, data_type)
        with ZipFile(BytesIO(content)) as archive:
            xml = archive.read("xl/worksheets/sheet1.xml")
            assert b"\r" not in xml
            assert b"&#13;" in xml or b"&#xD;" in xml
            tree = ET.fromstring(xml)
            cell = next(c for c in tree.findall(".//s:c", NS) if c.get("r") == "A1")
            assert "".join(t.text or "" for t in cell.findall(".//s:t", NS)) == payload
        assert ws["A1"].value == payload



def test_cr_save_adapter_preserves_other_archive_members_and_metadata():
    wb = Workbook()
    wb.active.append(["=1+1\r\n", "0012", 12])
    common.force_text_cells(wb.active)
    raw = BytesIO()
    # Explicit native class save supplies a real archive to the adapter unit
    # check; endpoint tests above also exercise the installed instance save.
    Workbook.save(wb, raw)
    target = BytesIO()
    common._save_preserving_carriage_returns(lambda stream: stream.write(raw.getvalue()), target)
    with ZipFile(BytesIO(raw.getvalue())) as before, ZipFile(BytesIO(target.getvalue())) as after:
        assert before.namelist() == after.namelist()
        assert before.comment == after.comment
        for member in before.infolist():
            name = member.filename
            expected = before.read(name)
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                expected = expected.replace(b"\r", b"&#13;")
            assert after.read(name) == expected
            saved = after.getinfo(name)
            assert (saved.compress_type, saved.date_time, saved.external_attr, saved.comment) == (
                member.compress_type, member.date_time, member.external_attr, member.comment,
            )
    loaded = assert_xlsx_contract(target.getvalue()).active
    assert_text(loaded["A1"], "=1+1\r\n")
    assert_text(loaded["B1"], "0012")
    assert (loaded["C1"].value, loaded["C1"].data_type) == (12, "n")


@pytest.mark.parametrize("grade", [None, "高一"])
def test_teacher_multisheet_seals_all_text_after_skipping_empty_course(grade):
    from app.services.xbk.exports.teacher_distribution import build_teacher_distribution_xlsx

    courses = [synthetic_row(course_code=code, teacher="教师\r\n=1+1") for code in ["0001", "0002", "0003"]]
    rows = [synthetic_row(class_name=cls, student_name=name)
            for cls, name in [("10班", "=1+1\r\n"), ("2班", "0012"), ("1班", "#N/A")]]
    db = CaptureDb(courses, rows, [], [synthetic_row(student_name="\r=1+1\r\n")])
    content = asyncio.run(build_teacher_distribution_xlsx(db, YEAR, TERM, None, None, None, grade)).getvalue()
    wb = assert_xlsx_contract(content)
    assert wb.sheetnames == ["0001", "0003"]
    assert [wb["0001"].cell(row, 2).value for row in range(6, 9)] == ["#N/A", "0012", "=1+1\r\n"]
    assert_text(wb["0003"]["B6"], "\r=1+1\r\n")
    for ws in wb:
        assert "教师\r\n=1+1" in ws["A2"].value
        assert {str(cell_range) for cell_range in ws.merged_cells.ranges} == {"A1:I1", "A2:I2", "A3:B4"}
        assert ws["B6"].data_type == "s" and ws["A2"].data_type == "s"
        assert ws.page_setup.orientation == "landscape"
    assert_text(wb["0001"]["A6"], "1班" if grade else "高一1班")
    assert len(db.calls) == 4  # Includes the empty course, not a fabricated sheet.


@pytest.mark.parametrize("scope", ["selections", "course_results"])
@pytest.mark.parametrize("payload", PAYLOADS)
def test_diagnostics_export_every_text_field(scope, payload):
    fields = [attr for _, attr in FIELDS[scope]]
    row = synthetic_row(**dict.fromkeys(fields, payload))
    wb = assert_xlsx_contract(ordinary_bytes(scope, None, diagnostic_rows=[row]))
    assert wb.sheetnames == ["data", "diagnostics"]
    assert wb["data"].max_row == 1
    assert [cell.value for cell in wb["data"][1]] == [header for header, _ in FIELDS[scope]]
    for col, (_, attr) in enumerate(FIELDS[scope], 1):
        assert_text(wb["diagnostics"].cell(2, col), getattr(row, attr))


@pytest.mark.parametrize("scope", ["selections", "course_results"])
def test_list_export_empty_sheets_keep_headers(scope):
    wb = assert_xlsx_contract(ordinary_bytes(scope, None))
    assert wb.sheetnames == ["data", "diagnostics"]
    for ws in wb:
        assert ws.max_row == 1
        assert [cell.value for cell in ws[1]] == [header for header, _ in FIELDS[scope]]
