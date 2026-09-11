"""Non-import XBK regressions; only an in-memory SQLite database is used.

The facade executes real SQL/transactions, never get_db or the configured engine.
PostgreSQL-specific sorting remains covered separately in an isolated test DB.
"""
import asyncio
import re
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import Session

from app.models import XbkCourse, XbkSelection, XbkStudent
from app.api.endpoints.xbk import analysis, bulk_ops, students, courses, selections
from app.api.endpoints.xbk._common import require_xbk_access
from app.schemas.xbk import XbkStudentUpsert, XbkCourseUpsert, XbkSelectionUpsert
from app.services.xbk.exports.class_distribution import build_class_distribution_xlsx
from app.services.xbk.exports.teacher_distribution import build_teacher_distribution_xlsx


DEFAULT_YEAR = "2026-2027"
OTHER_YEAR = "2025-2026"
FUTURE_YEAR = "2032-2033"


class MemoryDb:
    def __init__(self, session):
        self.session = session

    async def execute(self, stmt):
        return self.session.execute(stmt)

    def add(self, row):
        self.session.add(row)

    async def commit(self):
        self.session.commit()

    async def rollback(self):
        self.session.rollback()

    async def refresh(self, row):
        self.session.refresh(row)


@pytest.fixture
def db():
    engine = create_engine("sqlite:///:memory:")

    @event.listens_for(engine, "connect")
    def sqlite_regexp(connection, _):
        connection.create_function("regexp_replace", 4, lambda value, pattern, replacement, flags: re.sub(pattern, replacement, value))

    for model in (XbkStudent, XbkCourse, XbkSelection):
        model.__table__.create(engine)
    with Session(engine, expire_on_commit=False) as session:
        yield MemoryDb(session)
    engine.dispose()


def student(no="S1", year=2026, term="上学期", grade="高一", cls="1班", **extra):
    return XbkStudent(year=year, term=term, grade=grade, class_name=cls,
                      student_no=no, name=no, is_deleted=False, **extra)


def course(code="C1", year=2026, grade="高一"):
    return XbkCourse(year=year, term="上学期", grade=grade, course_code=code,
                     course_name=code, quota=3, is_deleted=False)


def selection(no="S1", code="C1", year=2026, term="上学期", grade="高一"):
    return XbkSelection(year=year, term=term, grade=grade, student_no=no,
                        name=no, course_code=code, is_deleted=False)


def seed(db, *rows):
    db.session.add_all(rows)
    db.session.commit()


def run_filtered(fn, db, **filters):
    params = dict(year=None, term=None, grade=None, class_name=None, db=db, _=None)
    params.update(filters)
    return asyncio.run(fn(**params))


@pytest.mark.parametrize("other", [dict(year=OTHER_YEAR), dict(term="下学期")])
def test_suspended_is_scoped_to_student_year_and_term(db, other):
    seed(db, student(), student(**other), course(), selection())
    summary = run_filtered(analysis.get_summary, db)
    assert summary["suspended_count"] == 1
    rows = run_filtered(analysis.students_without_selection, db)["items"]
    assert len(rows) == 1
    assert all(rows[0][key] == value for key, value in other.items())
    stats = run_filtered(analysis.course_stats, db)["items"]
    assert next(r for r in stats if r["course_code"] == "休学或其他")["count"] == 1


def test_stats_use_live_roster_grade_not_selection_snapshot(db):
    seed(db, student(), course(), selection(grade="高二"),
         student("S2"), selection("S2", code="", grade="高二"))
    summary = run_filtered(analysis.get_summary, db, grade="高一")
    assert summary["suspended_count"] == 0
    assert summary["selections"] == summary["unselected_count"] == 1
    assert run_filtered(analysis.students_without_selection, db, grade="高一")["items"] == []
    assert len(run_filtered(analysis.students_with_empty_selection, db, grade="高一")["items"]) == 1
    stats = run_filtered(analysis.course_stats, db, grade="高一")["items"]
    assert next(r for r in stats if r["course_code"] == "C1")["count"] == 1


def test_course_stats_exclude_deleted_and_orphan_students(db):
    gone = student("deleted")
    gone.is_deleted = True
    seed(db, student(), gone, course(), selection(), selection("deleted"), selection("orphan"))
    stats = run_filtered(analysis.course_stats, db)["items"]
    assert next(r for r in stats if r["course_code"] == "C1")["count"] == 1


def test_course_stats_use_catalog_and_keep_zero_enrollment_courses(db):
    deleted = course("DELETED")
    deleted.is_deleted = True
    seed(
        db,
        student(),
        course("C1"),
        course("ZERO"),
        deleted,
        selection(),
        selection(code="DELETED"),
        selection(code="ORPHAN"),
    )

    rows = run_filtered(
        analysis.course_stats,
        db,
        year=DEFAULT_YEAR,
        term="上学期",
        grade="高一",
    )["items"]
    real_rows = {
        row["course_code"]: row
        for row in rows
        if row["course_code"] not in {"未选", "休学或其他"}
    }

    assert set(real_rows) == {"C1", "ZERO"}
    assert real_rows["C1"]["count"] == 1
    assert real_rows["ZERO"]["count"] == 0
    assert all(row["grade"] == "高一" for row in real_rows.values())


def test_cross_grade_class_identity_and_course_capacity_use_course_grade(db):
    rows = [
        student(f"G1-{index}", grade="高一", cls=f"{index}班")
        for index in range(1, 19)
    ] + [
        student(f"G2-{index}", grade="高二", cls=f"{index}班")
        for index in range(1, 17)
    ]
    seed(db, *rows, course("G1-C", grade="高一"), course("G2-C", grade="高二"))

    course_rows = run_filtered(
        analysis.course_stats,
        db,
        year=DEFAULT_YEAR,
        term="上学期",
    )["items"]
    real_rows = {
        row["course_code"]: row
        for row in course_rows
        if row["course_code"] not in {"未选", "休学或其他"}
    }
    assert real_rows["G1-C"]["grade"] == "高一"
    assert real_rows["G1-C"]["class_count"] == 18
    assert real_rows["G1-C"]["allowed_total"] == 54
    assert real_rows["G2-C"]["grade"] == "高二"
    assert real_rows["G2-C"]["class_count"] == 16
    assert real_rows["G2-C"]["allowed_total"] == 48

    class_rows = run_filtered(
        analysis.class_stats,
        db,
        year=DEFAULT_YEAR,
        term="上学期",
    )["items"]
    assert len(class_rows) == 34
    assert len({(row["grade"], row["class_name"]) for row in class_rows}) == 34

    filtered_rows = run_filtered(
        analysis.course_stats,
        db,
        year=DEFAULT_YEAR,
        term="上学期",
        class_name="1班",
    )["items"]
    filtered_real_rows = {
        row["course_code"]: row
        for row in filtered_rows
        if row["course_code"] not in {"未选", "休学或其他"}
    }
    assert set(filtered_real_rows) == {"G1-C", "G2-C"}
    assert {row["class_count"] for row in filtered_real_rows.values()} == {1}
    assert run_filtered(
        analysis.get_summary,
        db,
        year=DEFAULT_YEAR,
        term="上学期",
        class_name="1班",
    )["courses"] == len(filtered_real_rows)


def test_cross_grade_class_filter_keeps_student_course_and_stats_scopes_aligned(db):
    seed(
        db,
        student("G1-S1", grade="高一", cls="1班"),
        student("G2-S1", grade="高二", cls="1班"),
        student("G2-S2", grade="高二", cls="2班"),
        course("G1-C", grade="高一"),
        course("G2-C", grade="高二"),
        selection("G1-S1", "G1-C", grade="高二"),
        selection("G2-S1", "G2-C", grade="高一"),
        selection("G2-S2", "G2-C", grade="高一"),
    )

    filters = dict(year=DEFAULT_YEAR, term="上学期", class_name="1班")
    summary = run_filtered(analysis.get_summary, db, **filters)
    stats = run_filtered(analysis.course_stats, db, **filters)["items"]
    real_rows = {
        row["course_code"]: row
        for row in stats
        if row["course_code"] not in {"未选", "休学或其他"}
    }
    assert summary == {
        "students": 2,
        "courses": 2,
        "selections": 2,
        "unselected_count": 0,
        "suspended_count": 0,
    }
    assert {code: row["count"] for code, row in real_rows.items()} == {"G1-C": 1, "G2-C": 1}
    assert {(row["grade"], row["class_count"]) for row in real_rows.values()} == {
        ("高一", 1),
        ("高二", 1),
    }

    grade_filters = {**filters, "grade": "高一"}
    grade_summary = run_filtered(analysis.get_summary, db, **grade_filters)
    grade_stats = run_filtered(analysis.course_stats, db, **grade_filters)["items"]
    grade_real_rows = [
        row for row in grade_stats if row["course_code"] not in {"未选", "休学或其他"}
    ]
    assert grade_summary["students"] == grade_summary["courses"] == grade_summary["selections"] == 1
    assert [(row["course_code"], row["grade"], row["count"]) for row in grade_real_rows] == [
        ("G1-C", "高一", 1)
    ]


def test_bulk_all_deletes_empty_and_orphan_records_but_not_other_term(db):
    seed(db, student(), course(), selection(code=""), selection("orphan", "missing"),
         selection(term="下学期"))
    result = asyncio.run(bulk_ops.delete_data("all", DEFAULT_YEAR, "上学期", None, None, db, {}))
    assert result["deleted"] == 4
    remaining = db.session.scalars(select(XbkSelection)).all()
    assert len(remaining) == 1 and remaining[0].term == "下学期"


@pytest.mark.parametrize("scope", ["all", "courses"])
def test_bulk_class_filter_cannot_delete_shared_courses(db, scope):
    seed(db, student(), student("S2", cls="2班"), course(), selection(), selection("S2"))
    with pytest.raises(HTTPException) as exc:
        asyncio.run(bulk_ops.delete_data(scope, DEFAULT_YEAR, "上学期", None, "1班", db, {}))
    assert exc.value.status_code == 400
    assert len(db.session.scalars(select(XbkSelection)).all()) == 2
    assert len(db.session.scalars(select(XbkCourse)).all()) == 1


@pytest.mark.parametrize("kind", ["student", "course", "selection"])
def test_update_unique_conflict_is_409_and_transaction_rolled_back(db, kind):
    s1, s2, c1, c2 = student(), student("S2"), course(), course("C2")
    a, b = selection(), selection(code="C2")
    seed(db, s1, s2, c1, c2, a, b)
    if kind == "student":
        payload = XbkStudentUpsert(year=2026, term="上学期", class_name="1班", student_no="S2", name="conflict")
        fn, pk = students.update_student, s1.id
    elif kind == "course":
        payload = XbkCourseUpsert(year=2026, term="上学期", course_code="C2", course_name="conflict")
        fn, pk = courses.update_course, c1.id
    else:
        payload = XbkSelectionUpsert(year=2026, term="上学期", student_no="S1", course_code="C2")
        fn, pk = selections.update_selection, a.id
    with pytest.raises(HTTPException) as exc:
        asyncio.run(fn(pk, payload, db, {}))
    assert exc.value.status_code == 409
    assert db.session.scalar(select(XbkStudent.name).where(XbkStudent.id == s1.id)) == "S1"


@pytest.mark.parametrize("builder,args", [
    (build_class_distribution_xlsx, (DEFAULT_YEAR, "上学期", None, None, None, None)),
    (build_teacher_distribution_xlsx, (DEFAULT_YEAR, "上学期", None, None, None)),
])
def test_empty_distribution_is_a_valid_workbook(builder, args):
    result = SimpleNamespace(all=lambda: [], scalars=lambda: SimpleNamespace(all=lambda: []))
    db = SimpleNamespace(execute=AsyncMock(return_value=result))
    output = asyncio.run(builder(db, *args))
    wb = load_workbook(output)
    assert wb.sheetnames
    assert wb.active["A1"].value == "当前筛选条件下暂无数据"


@pytest.mark.parametrize("enabled,user,allowed", [
    (False, None, False), (False, {"role_code": "teacher"}, False),
    (False, {"role_code": "student"}, False),
    (False, {"role_code": "admin"}, True), (False, {"role_code": "super_admin"}, True),
    (True, None, True), (True, {"role_code": "student"}, True),
])
def test_feature_flag_read_access_matrix(enabled, user, allowed):
    result = SimpleNamespace(scalar_one_or_none=lambda: SimpleNamespace(value={"enabled": enabled}))
    db = SimpleNamespace(execute=AsyncMock(return_value=result))
    if allowed:
        assert asyncio.run(require_xbk_access(db, user)) == user
    else:
        with pytest.raises(HTTPException) as exc:
            asyncio.run(require_xbk_access(db, user))
        assert exc.value.status_code == 403


def test_course_stats_include_empty_selection_without_catalog_sentinel(db):
    seed(db, student(), selection(code=""))
    rows = run_filtered(analysis.course_stats, db)["items"]
    assert len(rows) == 1 and rows[0]["course_code"] == "未选" and rows[0]["count"] == 1


@pytest.mark.parametrize("kind", ["student", "course"])
def test_soft_delete_cascades_and_recreate_does_not_revive_selections(db, kind):
    stu, crs, sel, other = student(), course(), selection(), selection(year=2025)
    seed(db, stu, crs, sel, other)
    if kind == "student":
        asyncio.run(students.delete_student(stu.id, db, {}))
        payload = XbkStudentUpsert(year=2026, term="上学期", class_name="1班", student_no="S1", name="S1")
        asyncio.run(students.create_student(payload, db, {}))
    else:
        asyncio.run(courses.delete_course(crs.id, db, {}))
        payload = XbkCourseUpsert(year=2026, term="上学期", course_code="C1", course_name="C1")
        asyncio.run(courses.create_course(payload, db, {}))
    db.session.expire_all()
    assert db.session.get(XbkSelection, sel.id).is_deleted is True
    assert db.session.get(XbkSelection, other.id).is_deleted is False


@pytest.mark.parametrize("scope", ["students", "courses"])
def test_bulk_parent_cascade_ignores_stale_selection_grade(db, scope):
    seed(db, student(), course(), selection(grade="高二"), student("S2", grade="高二"),
         course("C2", grade="高二"), selection("S2", "C2", grade="高二"))
    asyncio.run(bulk_ops.delete_data(scope, DEFAULT_YEAR, "上学期", "高一", None, db, {}))
    assert [row.student_no for row in db.session.scalars(select(XbkSelection)).all()] == ["S2"]


@pytest.mark.parametrize("scope", ["students", "selections"])
def test_bulk_class_scope_preserves_other_classes_and_courses(db, scope):
    seed(db, student(), student("S2", cls="2班"), course(), selection(), selection("S2"))
    asyncio.run(bulk_ops.delete_data(scope, DEFAULT_YEAR, "上学期", "高一", "1班", db, {}))
    assert [r.student_no for r in db.session.scalars(select(XbkSelection)).all()] == ["S2"]
    assert len(db.session.scalars(select(XbkCourse)).all()) == 1


def test_bulk_delete_rolls_back_on_later_failure(db):
    seed(db, student(), course(), selection())
    execute = db.execute
    calls = 0

    async def fail_second(stmt):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise RuntimeError("injected failure after selection deletion")
        return await execute(stmt)

    db.execute = fail_second
    with pytest.raises(RuntimeError):
        asyncio.run(bulk_ops.delete_data("all", DEFAULT_YEAR, "上学期", None, None, db, {}))
    assert len(db.session.scalars(select(XbkSelection)).all()) == 1


@pytest.mark.parametrize("endpoint", [courses.list_courses, selections.list_course_results])
def test_list_numeric_sort_accepts_full_identifier_width(endpoint):
    result = SimpleNamespace(scalar_one=lambda: 0, all=lambda: [])
    result.scalars = lambda: result
    db = SimpleNamespace(execute=AsyncMock(return_value=result))
    params = dict(year=DEFAULT_YEAR, term="上学期", grade=None, search_text=None, page=1, size=50, db=db, _=None)
    if endpoint is selections.list_course_results:
        params["class_name"] = None
    asyncio.run(endpoint(**params))
    assert "NUMERIC(50, 0)" in str(db.execute.call_args.args[0])


def test_bulk_all_grade_does_not_trust_stale_grade_over_live_roster(db):
    seed(db, student(), course(), selection(), student("S2", grade="高二"),
         course("C2", grade="高二"), selection("S2", "C2", grade="高一"),
         selection("orphan", "missing", grade="高一"))
    asyncio.run(bulk_ops.delete_data("all", DEFAULT_YEAR, "上学期", "高一", None, db, {}))
    assert [r.student_no for r in db.session.scalars(select(XbkSelection)).all()] == ["S2"]


def test_bulk_selection_grade_can_clean_orphan_records_without_other_grade(db):
    seed(db, selection("orphan1", grade="高一"), selection("orphan2", grade="高二"))
    asyncio.run(bulk_ops.delete_data("selections", DEFAULT_YEAR, "上学期", "高一", None, db, {}))
    assert [r.student_no for r in db.session.scalars(select(XbkSelection)).all()] == ["orphan2"]


@pytest.mark.parametrize("empty_code", ["", "未选"])
@pytest.mark.parametrize("view", ["summary", "empty-selection", "course-stats", "unselected-export"])
def test_unselected_aliases_agree_across_analysis_and_export(db, empty_code, view):
    """Legacy blanks and the import-normalized sentinel mean the same thing."""
    from io import BytesIO
    from app.api.endpoints.xbk import import_export

    seed(db, student("SELECTED", year=2032), student("EMPTY", year=2032),
         course("C1", year=2032), selection("SELECTED", "C1", year=2032),
         selection("EMPTY", empty_code, year=2032))
    filters = dict(year=FUTURE_YEAR, term="上学期", grade="高一", class_name="1班")
    if view == "summary":
        result = run_filtered(analysis.get_summary, db, **filters)
        assert result["students"] == 2
        assert result["selections"] == 1
        assert result["unselected_count"] == 1
        assert result["suspended_count"] == 0
    elif view == "empty-selection":
        result = run_filtered(analysis.students_with_empty_selection, db, **filters)
        assert [row["student_no"] for row in result["items"]] == ["EMPTY"]
    elif view == "course-stats":
        rows = run_filtered(analysis.course_stats, db, **filters)["items"]
        assert sum(row["count"] for row in rows if row["course_code"] in {"", "未选"}) == 1
        assert sum(row["count"] for row in rows if row["course_code"] == "C1") == 1
        assert sum(row["count"] for row in rows) == 2
    else:
        async def export_bytes():
            response = await import_export.export_data(
                scope="unselected", **filters, search_text=None, format="xlsx", db=db, _={})
            return b"".join([chunk async for chunk in response.body_iterator])
        sheet = load_workbook(BytesIO(asyncio.run(export_bytes()))).active
        rows = list(sheet.values)
        assert len(rows) == 2
        assert rows[1][rows[0].index("学号")] == "EMPTY"


def _export_records(db, scope, *, sheet="data", **filters):
    from io import BytesIO
    from app.api.endpoints.xbk import import_export

    async def content():
        response = await import_export.export_data(
            scope=scope,
            year=filters.get("year", FUTURE_YEAR),
            term=filters.get("term"),
            grade=filters.get("grade", "高一"),
            class_name=filters.get("class_name", "1班"),
            search_text=filters.get("search_text"),
            format="xlsx",
            db=db,
            _={},
        )
        return b"".join([chunk async for chunk in response.body_iterator])
    rows = list(load_workbook(BytesIO(asyncio.run(content())))[sheet].values)
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def test_course_stats_merge_legacy_blank_and_imported_unselected(db):
    seed(db, student("BLANK", year=2032), student("MARKER", year=2032),
         selection("BLANK", "", year=2032), selection("MARKER", "未选", year=2032))
    rows = run_filtered(analysis.course_stats, db, year=FUTURE_YEAR, term="上学期", grade="高一")["items"]
    assert len(rows) == 1
    assert rows[0]["course_code"] == rows[0]["course_name"] == "未选"
    assert rows[0]["count"] == 2


@pytest.mark.parametrize("scope", ["selections", "course_results"])
def test_selection_exports_use_live_roster_grade_and_keep_orphan_snapshot(db, scope):
    seed(
        db,
        student("IN", year=2032, grade="高一", cls="1班"),
        student("OUT", year=2032, grade="高二", cls="1班"),
        selection("IN", year=2032, grade="高二"),
        selection("OUT", year=2032, grade="高一"),
        selection("ORPHAN", year=2032, grade="高一"),
    )

    class_rows = _export_records(db, scope, term="上学期")
    assert [(row["学号"], row["年级"]) for row in class_rows] == [("IN", "高一")]

    grade_rows = _export_records(db, scope, term="上学期", class_name=None)
    assert [(row["学号"], row["年级"]) for row in grade_rows] == [("IN", "高一")]
    # XBK-07: preserve diagnostics explicitly, without mixing them into list rows.
    diagnostics = _export_records(db, scope, sheet="diagnostics", term="上学期", class_name=None)
    assert [(row["学号"], row["年级"]) for row in diagnostics] == [("ORPHAN", "高一")]
    assert _export_records(db, scope, sheet="diagnostics", term="上学期") == []


@pytest.mark.parametrize("empty_code", ["", "未选"])
def test_unselected_export_filters_live_student_grade(db, empty_code):
    seed(db, student("IN", year=2032), student("OUT", year=2032, grade="高二"),
         selection("IN", empty_code, year=2032, grade="高二"),
         selection("OUT", empty_code, year=2032, grade="高一"))
    assert [r["学号"] for r in _export_records(db, "unselected", term="上学期")] == ["IN"]


@pytest.mark.parametrize("term", [None, "上学期", "下学期"])
def test_suspended_export_correlates_semester_and_ignores_selection_grade(db, term):
    seed(db, student("SAME", year=2032), student("SAME", year=2032, term="下学期"),
         selection("SAME", "C1", year=2032, grade="高二"),
         student("OUT", year=2032, grade="高二"))
    rows = _export_records(db, "suspended", term=term)
    expected = [] if term == "上学期" else [(FUTURE_YEAR, "下学期", "SAME")]
    assert [(r["学年"], r["学期"], r["学号"]) for r in rows] == expected


@pytest.mark.parametrize("scope,removed,deleted", [
    ("students", {"S1", "S3"}, 4),
    ("courses", {"S2"}, 2),
    ("selections", {"S1", "S3", "orphan"}, 3),
    ("all", {"S1", "S2", "S3", "orphan"}, 7),
])
def test_bulk_grade_cascades_keep_persisted_identity_and_orphan_boundaries(db, scope, removed, deleted):
    deleted_parent = student("S3")
    deleted_parent.is_deleted = True
    seed(db, student(), student("S2", grade="高二"), deleted_parent,
         course(), course("C2", grade="高二"),
         selection("S1", "C2", grade="高二"),
         selection("S2", "C1", grade="高一"),
         selection("S3", "C2", grade="高二"),
         selection("orphan", "missing", grade="高一"),
         selection("other-orphan", "missing", grade="高二"),
         selection("other-year", "C1", year=2025))
    result = asyncio.run(bulk_ops.delete_data(scope, DEFAULT_YEAR, "上学期", "高一", None, db, {}))
    assert result == {"deleted": deleted}
    remaining = set(db.session.scalars(select(XbkSelection.student_no)).all())
    assert remaining == {"S1", "S2", "S3", "orphan", "other-orphan", "other-year"} - removed
    expected_students = {"S2"} if scope in {"all", "students"} else {"S1", "S2", "S3"}
    expected_courses = {"C2"} if scope in {"all", "courses"} else {"C1", "C2"}
    assert set(db.session.scalars(select(XbkStudent.student_no)).all()) == expected_students
    assert set(db.session.scalars(select(XbkCourse.course_code)).all()) == expected_courses


@pytest.mark.parametrize("failure", ["parent_delete", "commit"])
def test_bulk_failure_rolls_back_cascade_and_parents(db, failure, monkeypatch):
    seed(db, student(), course(), selection())
    original_execute = db.execute

    async def execute(stmt):
        if failure == "parent_delete" and stmt.is_delete and stmt.table.name == "xbk_students":
            raise RuntimeError("synthetic parent deletion failure")
        return await original_execute(stmt)

    monkeypatch.setattr(db, "execute", execute)
    if failure == "commit":
        monkeypatch.setattr(db, "commit", AsyncMock(side_effect=RuntimeError("synthetic commit failure")))
    rollback = AsyncMock(wraps=db.rollback)
    monkeypatch.setattr(db, "rollback", rollback)
    with pytest.raises(RuntimeError, match="synthetic"):
        asyncio.run(bulk_ops.delete_data("all", DEFAULT_YEAR, "上学期", None, None, db, {}))
    rollback.assert_awaited_once()
    with Session(db.session.bind) as fresh:
        for model in (XbkStudent, XbkCourse, XbkSelection):
            assert len(fresh.scalars(select(model)).all()) == 1
