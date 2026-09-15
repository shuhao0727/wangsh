"""Real ASGI/HTTP contract coverage for the trusted R3 workbook boundary."""

from __future__ import annotations

import asyncio
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook
import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.api.endpoints.xbk import course_selection_workbook as api
from app.core.deps import require_admin
from app.core.request_budget import RequestBudgetMiddleware
from app.db.database import get_db
from app.models import XbkCourse, XbkSelection, XbkStudent
from app.services.xbk.course_selection_workbook import (
    CATALOG_SHEET,
    INSTRUCTION,
    METADATA_SHEET,
    WorkbookCourse,
    WorkbookExpectation,
    WorkbookStudent,
    compute_baseline_id,
    write_metadata_sheet,
)
from app.services.xbk.course_selection_workbook_token import issue_course_selection_workbook_token
from app.services.xbk.selection_rules import ACTIVE_SELECTION_UNIQUE_CONSTRAINT


YEAR, TERM = "2038-2039", "上学期"


class MemoryDb:
    def __init__(self, session: Session):
        self.session = session
        self.commit_count = 0
        self.rollback_count = 0

    async def execute(self, statement):
        return self.session.execute(statement)

    def add(self, row):
        self.session.add(row)

    async def flush(self):
        self.session.flush()

    async def commit(self):
        self.commit_count += 1
        self.session.commit()

    async def rollback(self):
        self.rollback_count += 1
        self.session.rollback()


@pytest.fixture
def harness(monkeypatch):
    engine = create_engine("sqlite:///:memory:")
    for model in (XbkStudent, XbkCourse, XbkSelection):
        model.__table__.create(engine)
    session = Session(engine, expire_on_commit=False)
    db = MemoryDb(session)

    students = (
        WorkbookStudent("高一1班", "高一", "1班", "S1", "学生1", "C1"),
        WorkbookStudent("高一1班", "高一", "1班", "S2", "学生2", None),
    )
    courses = (
        WorkbookCourse("C1", "课程1", 5),
        WorkbookCourse("C2", "课程2", 5),
        WorkbookCourse("C3", "课程3", 5),
    )
    expected = WorkbookExpectation(
        year=YEAR,
        term=TERM,
        baseline_id=compute_baseline_id(YEAR, TERM, students, courses),
        students=students,
        courses=courses,
    )
    for student in students:
        session.add(XbkStudent(
            year=YEAR, term=TERM, grade=student.grade,
            class_name=student.class_name, student_no=student.student_no,
            name=student.name, is_deleted=False,
        ))
    for course in courses:
        session.add(XbkCourse(
            year=YEAR, term=TERM, grade="高一", course_code=course.course_code,
            course_name=course.course_name, quota=course.quota, is_deleted=False,
        ))
    session.add(XbkSelection(
        year=YEAR, term=TERM, grade="高一", student_no="S1",
        name="学生1", course_code="C1", is_deleted=False,
    ))
    session.commit()

    async def expectation(_db, year, term):
        assert _db is db
        assert (year, term) == (YEAR, TERM)
        return expected

    monkeypatch.setattr(api, "load_current_workbook_expectation", expectation)
    identity = {"id": 101, "role_code": "admin"}

    async def admin():
        return identity

    async def shared_db():
        yield db

    app = FastAPI()
    app.add_middleware(RequestBudgetMiddleware)
    app.include_router(api.router, prefix="/api/v1/xbk")
    app.dependency_overrides[require_admin] = admin
    app.dependency_overrides[get_db] = shared_db
    app.dependency_overrides[api.get_workbook_transaction_db] = shared_db

    yield {"app": app, "db": db, "identity": identity, "expected": expected, "session": session}
    session.close()
    engine.dispose()


def workbook_bytes(expected: WorkbookExpectation, submitted=(None, None)) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    catalog = workbook.create_sheet(CATALOG_SHEET)
    catalog.append(["课程目录"])
    catalog.append(["课程代码", "课程名称", "课程负责人", "各班限报人数", "上课地点"])
    for course in expected.courses:
        catalog.append([course.course_code, course.course_name, None, course.quota, None])
    catalog.protection.sheet = True

    sheets = {}
    for item, value in zip(expected.students, submitted):
        sheet = sheets.get(item.sheet)
        if sheet is None:
            sheet = workbook.create_sheet(item.sheet)
            sheets[item.sheet] = sheet
            sheet.append(["班级", "学号", "姓名", "课程代码", None, INSTRUCTION])
            sheet.cell(2, 6, "课程代码")
            sheet.cell(2, 7, "课程名称")
            sheet.cell(2, 8, "本班限额")
            sheet.cell(2, 9, "已选")
            sheet.cell(2, 10, "剩余")
            sheet.protection.sheet = True
        row = 2 + sum(
            1 for student in expected.students
            if student.sheet == item.sheet and student.student_no < item.student_no
        )
        sheet.cell(row, 1, item.class_name)
        sheet.cell(row, 2, item.student_no)
        sheet.cell(row, 3, item.name)
        sheet.cell(row, 4, value)

    metadata = workbook.create_sheet(METADATA_SHEET)
    write_metadata_sheet(
        metadata, year=expected.year, term=expected.term,
        students=expected.students, courses=expected.courses,
    )
    metadata.sheet_state = "veryHidden"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


async def request(app, method: str, path: str, **kwargs):
    async with AsyncClient(
        transport=ASGITransport(app=app, raise_app_exceptions=False),
        base_url="http://isolated.invalid",
    ) as client:
        return await client.request(method, path, **kwargs)


def upload(content: bytes, name="students.xlsx"):
    return {"file": (name, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def zip_bytes(entries: list[tuple[str, bytes]]) -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name, payload in entries:
            archive.writestr(name, payload)
    return output.getvalue()


def test_preview_and_confirm_use_signed_admin_bound_token_and_real_r3_apply(harness):
    content = workbook_bytes(harness["expected"], ("C2", None))

    async def scenario():
        preview = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/preview",
            data={"year": YEAR, "term": TERM}, files=upload(content),
        )
        assert preview.status_code == 200, preview.text
        body = preview.json()
        assert body["year"] == YEAR and body["term"] == TERM
        assert body["changed"] == 1 and body["unchanged"] == 1
        assert body["changed_rows"][0]["student_no"] == "S1"
        assert body["changed_rows"][0]["target_course"] == "C2"
        assert body["notice"].startswith("预览不预留名额")
        assert body["preview_token"].count(".") == 2

        confirm = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/confirm",
            data={"year": YEAR, "term": TERM, "preview_token": body["preview_token"]},
            files=upload(content),
        )
        assert confirm.status_code == 200, confirm.text
        assert confirm.json() == {
            "plan_id": body["plan_id"], "status": "applied",
            "changed": 1, "inserted": 0, "updated": 1,
        }

        repeated = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/confirm",
            data={"year": YEAR, "term": TERM, "preview_token": body["preview_token"]},
            files=upload(content),
        )
        assert repeated.status_code == 200, repeated.text
        assert repeated.json() == {
            "plan_id": body["plan_id"], "status": "already_applied",
            "changed": 0, "inserted": 0, "updated": 0,
        }

    asyncio.run(scenario())
    active = harness["session"].scalars(
        select(XbkSelection).where(XbkSelection.is_deleted.is_(False))
    ).all()
    assert [(row.student_no, row.course_code) for row in active] == [("S1", "C2")]
    assert harness["db"].commit_count == 2


def test_confirmation_reauth_rejects_another_admin_before_apply(harness):
    content = workbook_bytes(harness["expected"], ("C2", None))

    async def scenario():
        preview = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/preview",
            data={"year": YEAR, "term": TERM}, files=upload(content),
        )
        token = preview.json()["preview_token"]
        harness["identity"]["id"] = 202
        confirm = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/confirm",
            data={"year": YEAR, "term": TERM, "preview_token": token}, files=upload(content),
        )
        assert confirm.status_code == 403
        assert confirm.json()["detail"]["code"] == "XBK_WORKBOOK_PREVIEW_TOKEN_IDENTITY_MISMATCH"

    asyncio.run(scenario())
    assert harness["db"].commit_count == 0


def test_tampered_preview_token_returns_structured_401(harness):
    content = workbook_bytes(harness["expected"], ("C2", None))

    async def scenario():
        preview = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/preview",
            data={"year": YEAR, "term": TERM}, files=upload(content),
        )
        token = preview.json()["preview_token"]
        replacement = "A" if token[-1] != "A" else "B"
        confirm = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/confirm",
            data={"year": YEAR, "term": TERM, "preview_token": token[:-1] + replacement},
            files=upload(content),
        )
        assert confirm.status_code == 401
        assert confirm.json()["detail"]["code"] == "XBK_WORKBOOK_PREVIEW_TOKEN_INVALID"

    asyncio.run(scenario())
    assert harness["db"].commit_count == 0


def test_confirmation_rejects_different_valid_workbook_plan(harness):
    preview_content = workbook_bytes(harness["expected"], ("C2", None))
    changed_content = workbook_bytes(harness["expected"], ("C3", None))

    async def scenario():
        preview = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/preview",
            data={"year": YEAR, "term": TERM}, files=upload(preview_content),
        )
        confirm = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/confirm",
            data={"year": YEAR, "term": TERM, "preview_token": preview.json()["preview_token"]},
            files=upload(changed_content),
        )
        assert confirm.status_code == 409
        assert confirm.json()["detail"]["code"] == "XBK_WORKBOOK_PREVIEW_TOKEN_PLAN_MISMATCH"

    asyncio.run(scenario())
    assert harness["db"].commit_count == 0
    assert harness["db"].rollback_count == 1


def test_confirmation_rejects_course_quota_change_after_preview(harness):
    content = workbook_bytes(harness["expected"], ("C2", None))

    async def scenario():
        preview = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/preview",
            data={"year": YEAR, "term": TERM}, files=upload(content),
        )
        assert preview.status_code == 200, preview.text
        course = harness["session"].scalar(
            select(XbkCourse).where(XbkCourse.course_code == "C2")
        )
        course.quota += 1
        harness["session"].commit()

        confirm = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/confirm",
            data={
                "year": YEAR,
                "term": TERM,
                "preview_token": preview.json()["preview_token"],
            },
            files=upload(content),
        )
        assert confirm.status_code == 409, confirm.text
        assert confirm.json()["detail"]["code"] == "XBK_WORKBOOK_STALE_BASELINE"

    asyncio.run(scenario())
    assert harness["db"].commit_count == 0
    assert harness["db"].rollback_count == 1


def test_expired_token_is_rejected_before_workbook_or_database_processing(harness):
    token, _ = issue_course_selection_workbook_token(
        admin_id=101, year=YEAR, term=TERM, plan_id="a" * 64, now=1,
    )

    async def scenario():
        response = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/confirm",
            data={"year": YEAR, "term": TERM, "preview_token": token},
            files=upload(b"not-even-opened"),
        )
        assert response.status_code == 401
        assert response.json()["detail"]["code"] == "XBK_WORKBOOK_PREVIEW_TOKEN_EXPIRED"

    asyncio.run(scenario())
    assert harness["db"].commit_count == 0


def test_malformed_workbook_returns_structured_422(harness):
    async def scenario():
        response = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/preview",
            data={"year": YEAR, "term": TERM}, files=upload(b"not-an-xlsx"),
        )
        assert response.status_code == 422
        detail = response.json()["detail"]
        assert detail["code"] == "XBK_WORKBOOK_ARCHIVE_INVALID"
        assert detail["issues"] == []

    asyncio.run(scenario())


def test_archive_budget_returns_structured_413_before_openpyxl(harness, monkeypatch):
    monkeypatch.setattr(
        "app.services.xbk.course_selection_workbook.MAX_XLSX_COMPRESSION_RATIO", 2
    )
    content = zip_bytes([("xl/worksheets/sheet1.xml", b"A" * 4096)])

    async def scenario():
        response = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/preview",
            data={"year": YEAR, "term": TERM}, files=upload(content),
        )
        assert response.status_code == 413
        assert response.json()["detail"]["code"] == "XBK_WORKBOOK_ARCHIVE_TOO_LARGE"
        assert response.json()["detail"]["issues"] == []

    asyncio.run(scenario())


def test_zip_slip_returns_structured_422_before_openpyxl(harness):
    content = zip_bytes([("../escape.xml", b"x")])

    async def scenario():
        response = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/preview",
            data={"year": YEAR, "term": TERM}, files=upload(content),
        )
        assert response.status_code == 422
        assert response.json()["detail"]["code"] == "XBK_WORKBOOK_ARCHIVE_INVALID"
        assert response.json()["detail"]["issues"] == []

    asyncio.run(scenario())


def test_confirm_maps_active_selection_unique_conflict_to_structured_409(harness, monkeypatch):
    content = workbook_bytes(harness["expected"], ("C2", None))

    class UniqueViolation(Exception):
        diag = type("Diag", (), {"constraint_name": ACTIVE_SELECTION_UNIQUE_CONSTRAINT})()

    async def scenario():
        preview = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/preview",
            data={"year": YEAR, "term": TERM}, files=upload(content),
        )
        assert preview.status_code == 200, preview.text

        async def conflict(*_args, **_kwargs):
            raise IntegrityError("flush", {}, UniqueViolation("duplicate active selection"))

        monkeypatch.setattr(api, "apply_course_selection_workbook_plan", conflict)
        response = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/confirm",
            data={
                "year": YEAR,
                "term": TERM,
                "preview_token": preview.json()["preview_token"],
            },
            files=upload(content),
        )
        assert response.status_code == 409
        assert response.json()["detail"] == {
            "code": "XBK_WORKBOOK_ACTIVE_SELECTION_CONFLICT",
            "message": "学生选课已被并发修改，请重新预览后重试",
            "issues": [],
        }

    asyncio.run(scenario())
    assert harness["db"].rollback_count == 1


def test_global_request_budget_rejects_oversized_multipart_before_route(harness):
    async def scenario():
        response = await request(
            harness["app"], "POST", "/api/v1/xbk/course-selection-workbook/preview",
            data={"year": YEAR, "term": TERM},
            files=upload(b"x" * (2 * 1024 * 1024), name="oversized.xlsx"),
        )
        assert response.status_code == 413
        assert response.json() == {"detail": "请求体超过该接口允许预算"}

    asyncio.run(scenario())


def test_routes_require_admin_and_confirmation_uses_separate_transaction_dependency():
    routes = {route.path: route for route in api.router.routes}
    preview = routes["/course-selection-workbook/preview"]
    confirm = routes["/course-selection-workbook/confirm"]
    preview_dependencies = [item.call for item in preview.dependant.dependencies]
    confirm_dependencies = [item.call for item in confirm.dependant.dependencies]
    assert require_admin in preview_dependencies
    assert require_admin in confirm_dependencies
    assert get_db in preview_dependencies
    assert api.get_workbook_transaction_db in confirm_dependencies
    assert get_db not in confirm_dependencies
