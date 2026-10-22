"""Workbook/SQL-contract tests with fake query results; no configured DB access."""
import asyncio
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.api.endpoints.xbk import exports
from app.models import XbkCourse, XbkSelection, XbkStudent
from app.services.xbk.exports import course_selection, class_distribution, teacher_distribution


class Rows:
    def __init__(self, rows):
        self.rows = rows

    def all(self):
        return self.rows

    def scalars(self):
        return self


class CaptureDb:
    def __init__(self, *rows):
        self.results = iter(rows)
        self.statements = []

    async def execute(self, stmt):
        self.statements.append(stmt)
        return Rows(next(self.results))


def catalog(code="C1"):
    return XbkCourse(year=2026, term="上学期", grade="高一", course_code=code,
                     course_name="课程", teacher="老师", quota=3, location="A101", is_deleted=False)


def roster():
    return XbkStudent(year=2026, term="上学期", grade="高一", class_name="1班",
                      student_no="00001", name="学生", is_deleted=False)


def selected(student_no="00001", code="C1"):
    return XbkSelection(year=2026, term="上学期", student_no=student_no,
                        course_code=code, name="学生", is_deleted=False)


@pytest.mark.parametrize("kind,builder_name", [
    ("course-selection", "build_student_course_selection_xlsx"),
    ("distribution", "build_class_distribution_xlsx"),
    ("teacher-distribution", "build_teacher_distribution_xlsx"),
])
def test_export_endpoint_passes_grade_to_every_builder(monkeypatch, kind, builder_name):
    builder = AsyncMock(return_value=iter([b"xlsx"]))
    monkeypatch.setattr(exports, builder_name, builder)
    asyncio.run(exports.export_tables(kind, 2026, "上学期", "高一", None, None, None, object(), {}))
    assert builder.call_args.kwargs.get("grade") == "高一" or "高一" in builder.call_args.args


def test_student_template_honors_page_filters_for_roster_scope():
    # P0-2：年级/班级筛选现在真正作用于名册；课程目录仍按学年/学期整体加载。
    # 第三个结果集为空：该学期无任何名册时仍允许导出仅含目录的模板。
    db = CaptureDb([], [], [])
    asyncio.run(course_selection.build_student_course_selection_xlsx(
        db, 2026, "上学期", "1班", None, None, grade="高一"))
    assert "xbk_courses.grade =" not in str(db.statements[0])
    student_sql = str(db.statements[1])
    assert "xbk_students.grade =" in student_sql
    assert "xbk_students.class_name =" in student_sql


def test_teacher_export_filters_roster_by_grade_and_excludes_deleted_students():
    db = CaptureDb([catalog()], [])
    asyncio.run(teacher_distribution.build_teacher_distribution_xlsx(
        db, 2026, "上学期", None, None, None, grade="高一"))
    assert "xbk_courses.grade" in str(db.statements[0])
    sql = str(db.statements[1])
    assert "xbk_students.grade" in sql
    assert "LEFT OUTER JOIN xbk_students" not in sql


def test_class_export_uses_roster_grade_and_excludes_deleted_students():
    db = CaptureDb([])
    asyncio.run(class_distribution.build_class_distribution_xlsx(
        db, 2026, "上学期", "高一", None, None, None))
    sql = str(db.statements[0])
    assert "xbk_students.grade" in sql
    assert "LEFT OUTER JOIN xbk_students" not in sql


@pytest.mark.parametrize("code", ["001", "1234567890123456789012345", "H1-AI-02", "1.5"])
def test_student_template_preserves_course_identifiers_and_supports_text_codes(code):
    db = CaptureDb([catalog(code)], [roster()], [])
    output = asyncio.run(course_selection.build_student_course_selection_xlsx(
        db, 2026, "上学期", None, None, None))
    wb = load_workbook(output)
    assert wb["校本课程目录"]["A3"].value == code
    assert wb["高一1班"]["B2"].value == "00001"
    validation = next(iter(wb["高一1班"].data_validations.dataValidation))
    assert "ISNUMBER(D2)" not in validation.formula1
    assert 'VLOOKUP' in validation.formula1
    assert ",4,0)" in validation.formula1


def test_student_course_selection_populates_current_codes_by_class():
    second = roster()
    second.student_no = "00002"
    second.name = "未选学生"
    db = CaptureDb(
        [catalog("C1"), catalog("C2")],
        [roster(), second],
        [selected("00001", "C1")],
    )

    output = asyncio.run(course_selection.build_student_course_selection_xlsx(
        db, 2026, "上学期", None, None, None, grade="高一"
    ))
    ws = load_workbook(output)["高一1班"]

    assert ws["D2"].value == "C1"
    assert ws["D3"].value is None
    # A filtered export still uses the roster scope and never leaks an orphan
    # selection into another class sheet.
    assert len(db.statements) == 3
    assert "xbk_selections" in str(db.statements[2])
    assert "IN" in str(db.statements[2])


@pytest.mark.parametrize("builder,args", [
    (course_selection.build_student_course_selection_xlsx, (2026, "上学期", None, None, None)),
    (class_distribution.build_class_distribution_xlsx, (2026, "上学期", None, None, None, None)),
    (teacher_distribution.build_teacher_distribution_xlsx, (2026, "上学期", None, None, None)),
])
def test_export_numeric_sort_accepts_full_50_character_identifiers(builder, args):
    # course-selection 的第三条查询是既有选课快照；补齐结果集避免迭代耗尽。
    results = [[], [], []] if builder is course_selection.build_student_course_selection_xlsx else [[], []]
    db = CaptureDb(*results)
    asyncio.run(builder(db, *args))
    assert "NUMERIC(50, 0)" in str(db.statements[0])


# --- P0-2：学生选课表导出必须真正应用年级/班级筛选 ---------------------------
def student(no, grade, class_name):
    return XbkStudent(year=2026, term="上学期", grade=grade, class_name=class_name,
                      student_no=no, name=no, is_deleted=False)


def test_course_selection_export_only_builds_matching_class_sheet():
    selected = student("00002", "高二", "2班")
    db = CaptureDb([catalog("C1")], [selected], [])

    output = asyncio.run(course_selection.build_student_course_selection_xlsx(
        db, 2026, "上学期", "2班", None, None, grade="高二",
    ))

    workbook = load_workbook(output)
    class_sheets = [
        name for name in workbook.sheetnames
        if name not in {course_selection.CATALOG_SHEET, course_selection.METADATA_SHEET}
    ]
    assert class_sheets == ["高二2班"]
    assert workbook["高二2班"]["B2"].value == "00002"
    # 筛选未命中全校其他班级：名册查询带上年级/班级条件。
    assert "xbk_students.grade =" in str(db.statements[1])
    assert "xbk_students.class_name =" in str(db.statements[1])


@pytest.mark.parametrize("filters", [{"grade": "高三"}, {"class_name": "9班"}])
def test_course_selection_export_without_matching_class_returns_422(filters):
    # 该学期确实有名册，但筛选后没有匹配班级：必须明确报错，不能给空工作簿。
    db = CaptureDb([catalog("C1")], [], [1])

    with pytest.raises(HTTPException) as exc:
        asyncio.run(exports.export_tables(
            "course-selection", 2026, "上学期", filters.get("grade"),
            filters.get("class_name"), None, None, db, {},
        ))

    assert exc.value.status_code == 422
    assert "没有有效学生名册" in str(exc.value.detail)


def test_course_selection_export_keeps_empty_period_template():
    # 整个学期没有名册时仍返回仅含课程目录的可填写模板（R3 基线合同）。
    db = CaptureDb([catalog("C1")], [], [])
    output = asyncio.run(course_selection.build_student_course_selection_xlsx(
        db, 2026, "上学期", None, None, None,
    ))
    workbook = load_workbook(output)
    assert course_selection.CATALOG_SHEET in workbook.sheetnames
