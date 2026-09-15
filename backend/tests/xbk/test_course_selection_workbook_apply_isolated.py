"""Isolated R3 preview/confirm coverage without routers, network, or normal DB."""

import asyncio
from dataclasses import replace
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.models import XbkCourse, XbkSelection, XbkStudent
from app.services.xbk import course_selection_workbook_apply as service
from app.services.xbk.course_selection_workbook import (
    CATALOG_SHEET,
    INSTRUCTION,
    METADATA_SHEET,
    WorkbookCourse,
    WorkbookExpectation,
    WorkbookStudent,
    compute_baseline_id,
    write_metadata_sheet,
)
from app.services.xbk.course_selection_workbook_apply import (
    CourseSelectionWorkbookServiceError,
    apply_course_selection_workbook_plan,
    preview_course_selection_workbook,
)

YEAR, TERM = "2037-2038", "上学期"


class MemoryDb:
    def __init__(self, session):
        self.session = session
        self.rollback_count = 0

    async def execute(self, statement):
        return self.session.execute(statement)

    def add(self, row):
        self.session.add(row)

    async def flush(self):
        self.session.flush()

    async def rollback(self):
        self.rollback_count += 1
        self.session.rollback()

    async def commit(self):
        self.session.commit()


@pytest.fixture
def db():
    engine = create_engine("sqlite:///:memory:")
    for model in (XbkStudent, XbkCourse, XbkSelection):
        model.__table__.create(engine)
    with Session(engine, expire_on_commit=False) as session:
        yield MemoryDb(session)
    engine.dispose()


def expectation(*, baselines=("C1", None), quota=5):
    students = (
        WorkbookStudent("高一1班", "高一", "1班", "S1", "学生1", baselines[0]),
        WorkbookStudent("高一1班", "高一", "1班", "S2", "学生2", baselines[1]),
    )
    courses = (
        WorkbookCourse("C1", "课程1", quota),
        WorkbookCourse("C2", "课程2", quota),
        WorkbookCourse("C3", "课程3", quota),
    )
    return WorkbookExpectation(
        year=YEAR,
        term=TERM,
        baseline_id=compute_baseline_id(YEAR, TERM, students, courses),
        students=students,
        courses=courses,
    )


def workbook_bytes(expected, submitted=(None, None)):
    wb = Workbook()
    wb.remove(wb.active)

    catalog = wb.create_sheet(CATALOG_SHEET)
    catalog.append(["课程目录"])
    catalog.append(["课程代码", "课程名称", "课程负责人", "各班限报人数", "上课地点"])
    for course in expected.courses:
        catalog.append([course.course_code, course.course_name, None, course.quota, None])
    catalog.protection.sheet = True

    sheets = {}
    for item, value in zip(expected.students, submitted):
        ws = sheets.get(item.sheet)
        if ws is None:
            ws = wb.create_sheet(item.sheet)
            sheets[item.sheet] = ws
            ws.append(["班级", "学号", "姓名", "课程代码", None, INSTRUCTION])
            ws.cell(2, 6, "课程代码")
            ws.cell(2, 7, "课程名称")
            ws.cell(2, 8, "本班限额")
            ws.cell(2, 9, "已选")
            ws.cell(2, 10, "剩余")
            ws.protection.sheet = True
        row = 2 + sum(1 for student in expected.students if student.sheet == item.sheet and student.student_no < item.student_no)
        ws.cell(row, 1, item.class_name)
        ws.cell(row, 2, item.student_no)
        ws.cell(row, 3, item.name)
        ws.cell(row, 4, value)

    metadata = wb.create_sheet(METADATA_SHEET)
    write_metadata_sheet(
        metadata,
        year=expected.year,
        term=expected.term,
        students=expected.students,
        courses=expected.courses,
    )
    metadata.sheet_state = "veryHidden"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def seed(db, expected, *, selections=()):
    for student in expected.students:
        db.session.add(XbkStudent(
            year=YEAR,
            term=TERM,
            grade=student.grade,
            class_name=student.class_name,
            student_no=student.student_no,
            name=student.name,
            is_deleted=False,
        ))
    for course in expected.courses:
        db.session.add(XbkCourse(
            year=YEAR,
            term=TERM,
            grade="高一",
            course_code=course.course_code,
            course_name=course.course_name,
            quota=course.quota,
            is_deleted=False,
        ))
    for student_no, code in selections:
        db.session.add(XbkSelection(
            year=YEAR,
            term=TERM,
            grade="高一",
            student_no=student_no,
            name=student_no,
            course_code=code,
            is_deleted=False,
        ))
    db.session.commit()


def active(db):
    return {
        row.student_no: row.course_code
        for row in db.session.scalars(select(XbkSelection).where(~XbkSelection.is_deleted)).all()
    }


def test_preview_uses_server_expectation_and_builds_stable_immutable_plan():
    expected = expectation()
    source = workbook_bytes(expected, ("C2", "未选"))

    first = preview_course_selection_workbook(source, expectation=expected)
    second = preview_course_selection_workbook(source, expectation=expected)

    assert first == second
    assert first.plan_id and len(first.plan_id) == 64
    assert [(row.action, row.baseline_course, row.target_course) for row in first.rows] == [
        ("select", "C1", "C2"),
        ("set_unselected", None, None),
    ]
    assert isinstance(first.rows, tuple)


def test_blank_means_no_change_while_explicit_unselected_means_cancel():
    expected = expectation(baselines=("C1", "C2"))
    plan = preview_course_selection_workbook(
        workbook_bytes(expected, (None, "未选")), expectation=expected,
    )

    assert plan.rows[0].action == "no_change"
    assert plan.rows[0].target_course == "C1"
    assert plan.rows[1].action == "set_unselected"
    assert plan.rows[1].target_course is None
    assert [row.student_no for row in plan.changed_rows] == ["S2"]


def test_invalid_workbook_returns_structured_parser_issues():
    expected = expectation()
    bad = workbook_bytes(expected, ("NOT-A-COURSE", None))

    with pytest.raises(CourseSelectionWorkbookServiceError) as caught:
        preview_course_selection_workbook(bad, expectation=expected)

    payload = caught.value.to_dict()
    assert payload["code"] == "XBK_WORKBOOK_INVALID"
    assert payload["issues"][0]["sheet"] == "高一1班"
    assert payload["issues"][0]["row"] == 2
    assert payload["issues"][0]["student_no"] == "S1"
    assert payload["issues"][0]["course_code"] == "NOT-A-COURSE"


def test_confirm_calls_shared_mutation_service_once_for_whole_changed_batch(db, monkeypatch):
    expected = expectation(baselines=("C1", "C3"))
    seed(db, expected, selections=(("S1", "C1"), ("S2", "C3")))
    plan = preview_course_selection_workbook(
        workbook_bytes(expected, ("C2", "未选")), expectation=expected,
    )
    original = service.apply_selection_mutations
    calls = []

    async def spy(session, mutations):
        calls.append(tuple(mutations))
        return await original(session, mutations)

    monkeypatch.setattr(service, "apply_selection_mutations", spy)
    result = asyncio.run(apply_course_selection_workbook_plan(
        db, plan, expected_plan_id=plan.plan_id, source_expectation=expected,
    ))

    assert result.status == "applied"
    assert result.changed == 2
    assert len(calls) == 1 and len(calls[0]) == 2
    assert all(item.replace_existing for item in calls[0])
    assert active(db) == {"S1": "C2", "S2": "未选"}


def test_repeat_confirmation_is_idempotent_and_does_not_call_writer_again(db, monkeypatch):
    expected = expectation()
    seed(db, expected, selections=(("S1", "C1"),))
    plan = preview_course_selection_workbook(
        workbook_bytes(expected, ("C2", None)), expectation=expected,
    )
    first = asyncio.run(apply_course_selection_workbook_plan(
        db, plan, expected_plan_id=plan.plan_id, source_expectation=expected,
    ))
    assert first.status == "applied"

    async def forbidden(*_args, **_kwargs):
        raise AssertionError("idempotent confirmation must not write")

    monkeypatch.setattr(service, "apply_selection_mutations", forbidden)
    second = asyncio.run(apply_course_selection_workbook_plan(
        db, plan, expected_plan_id=plan.plan_id, source_expectation=expected,
    ))
    assert second.status == "already_applied"
    assert active(db) == {"S1": "C2"}


def test_third_state_and_partially_applied_batch_are_rejected(db):
    expected = expectation(baselines=("C1", "C3"))
    seed(db, expected, selections=(("S1", "C3"), ("S2", "C3")))
    plan = preview_course_selection_workbook(
        workbook_bytes(expected, ("C2", "C1")), expectation=expected,
    )

    with pytest.raises(CourseSelectionWorkbookServiceError) as caught:
        asyncio.run(apply_course_selection_workbook_plan(
            db, plan, expected_plan_id=plan.plan_id, source_expectation=expected,
        ))

    assert caught.value.code == "XBK_WORKBOOK_STALE_BASELINE"
    assert caught.value.issues[0].student_no == "S1"
    assert active(db) == {"S1": "C3", "S2": "C3"}
    assert db.rollback_count == 1


def test_tampered_plan_and_changed_expectation_are_rejected_before_writes(db):
    expected = expectation()
    plan = preview_course_selection_workbook(
        workbook_bytes(expected, ("C2", None)), expectation=expected,
    )
    tampered_row = replace(plan.rows[0], target_course="C3")
    tampered = replace(plan, rows=(tampered_row, *plan.rows[1:]))

    with pytest.raises(CourseSelectionWorkbookServiceError) as caught:
        asyncio.run(apply_course_selection_workbook_plan(
            db, tampered, expected_plan_id=plan.plan_id, source_expectation=expected,
        ))
    assert caught.value.code == "XBK_WORKBOOK_PLAN_TAMPERED"

    changed = expectation(quota=6)
    with pytest.raises(CourseSelectionWorkbookServiceError) as caught:
        asyncio.run(apply_course_selection_workbook_plan(
            db, plan, expected_plan_id=plan.plan_id, source_expectation=changed,
        ))
    assert caught.value.code == "XBK_WORKBOOK_STALE_BASELINE"


def test_no_change_plan_is_safe_noop_after_locked_baseline_check(db, monkeypatch):
    expected = expectation()
    seed(db, expected, selections=(("S1", "C1"),))
    plan = preview_course_selection_workbook(
        workbook_bytes(expected), expectation=expected,
    )

    async def forbidden(*_args, **_kwargs):
        raise AssertionError("no-change plan must not write")

    monkeypatch.setattr(service, "apply_selection_mutations", forbidden)
    result = asyncio.run(apply_course_selection_workbook_plan(
        db, plan, expected_plan_id=plan.plan_id, source_expectation=expected,
    ))
    assert result.status == "no_changes"


def test_writer_conflict_rolls_back_entire_batch_and_returns_structured_error(db, monkeypatch):
    expected = expectation(baselines=("C1", "C3"))
    seed(db, expected, selections=(("S1", "C1"), ("S2", "C3")))
    plan = preview_course_selection_workbook(
        workbook_bytes(expected, ("C2", "C1")), expectation=expected,
    )

    async def fail(session, mutations):
        session.session.get(XbkSelection, 1).course_code = "BROKEN-PARTIAL"
        await session.flush()
        raise HTTPException(status_code=409, detail="课程已满")

    monkeypatch.setattr(service, "apply_selection_mutations", fail)
    with pytest.raises(CourseSelectionWorkbookServiceError) as caught:
        asyncio.run(apply_course_selection_workbook_plan(
            db, plan, expected_plan_id=plan.plan_id, source_expectation=expected,
        ))

    assert caught.value.code == "XBK_WORKBOOK_APPLY_CONFLICT"
    assert caught.value.issues[0].message == "课程已满"
    assert active(db) == {"S1": "C1", "S2": "C3"}
    assert db.rollback_count == 1


def test_cancellation_rolls_back_and_propagates_cancelled_error(db, monkeypatch):
    expected = expectation()
    seed(db, expected, selections=(("S1", "C1"),))
    plan = preview_course_selection_workbook(
        workbook_bytes(expected, ("C2", None)), expectation=expected,
    )

    async def cancel(session, mutations):
        session.session.get(XbkSelection, 1).course_code = "BROKEN-PARTIAL"
        await session.flush()
        raise asyncio.CancelledError()

    monkeypatch.setattr(service, "apply_selection_mutations", cancel)
    with pytest.raises(asyncio.CancelledError):
        asyncio.run(apply_course_selection_workbook_plan(
            db, plan, expected_plan_id=plan.plan_id, source_expectation=expected,
        ))

    assert active(db) == {"S1": "C1"}
    assert db.rollback_count == 1


def test_no_change_plan_rejects_selection_that_moved_after_export(db):
    expected = expectation()
    seed(db, expected, selections=(("S1", "C2"),))
    plan = preview_course_selection_workbook(
        workbook_bytes(expected), expectation=expected,
    )

    with pytest.raises(CourseSelectionWorkbookServiceError) as caught:
        asyncio.run(apply_course_selection_workbook_plan(
            db, plan, expected_plan_id=plan.plan_id, source_expectation=expected,
        ))

    assert caught.value.code == "XBK_WORKBOOK_STALE_BASELINE"
    assert caught.value.issues[0].student_no == "S1"
    assert db.rollback_count == 1


def test_recomputed_digest_cannot_authorize_plan_outside_server_contract(db):
    expected = expectation()
    plan = preview_course_selection_workbook(
        workbook_bytes(expected, ("C2", None)), expectation=expected,
    )
    forged_row = replace(plan.rows[0], name="伪造姓名")
    unsigned_forged = replace(plan, rows=(forged_row, *plan.rows[1:]), plan_id="")
    forged_id = service._plan_digest(service._plan_payload(unsigned_forged))
    forged = replace(unsigned_forged, plan_id=forged_id)

    with pytest.raises(CourseSelectionWorkbookServiceError) as caught:
        asyncio.run(apply_course_selection_workbook_plan(
            db, forged, expected_plan_id=forged_id, source_expectation=expected,
        ))

    assert caught.value.code == "XBK_WORKBOOK_PLAN_TAMPERED"


def test_new_active_student_outside_export_baseline_rejects_whole_batch(db, monkeypatch):
    expected = expectation()
    seed(db, expected, selections=(("S1", "C1"),))
    db.session.add(XbkStudent(
        year=YEAR,
        term=TERM,
        grade="高一",
        class_name="1班",
        student_no="S3",
        name="学生3",
        is_deleted=False,
    ))
    db.session.commit()
    plan = preview_course_selection_workbook(
        workbook_bytes(expected, ("C2", None)), expectation=expected,
    )

    async def forbidden(*_args, **_kwargs):
        raise AssertionError("stale roster must block the whole batch")

    monkeypatch.setattr(service, "apply_selection_mutations", forbidden)
    with pytest.raises(CourseSelectionWorkbookServiceError) as caught:
        asyncio.run(apply_course_selection_workbook_plan(
            db, plan, expected_plan_id=plan.plan_id, source_expectation=expected,
        ))

    assert caught.value.code == "XBK_WORKBOOK_STALE_BASELINE"
    assert [(issue.code, issue.student_no) for issue in caught.value.issues] == [
        ("XBK_WORKBOOK_STALE_STUDENT", "S3"),
    ]
    assert active(db) == {"S1": "C1"}
    assert db.rollback_count == 1


def test_new_active_course_outside_export_baseline_rejects_whole_batch(db, monkeypatch):
    expected = expectation()
    seed(db, expected, selections=(("S1", "C1"),))
    db.session.add(XbkCourse(
        year=YEAR,
        term=TERM,
        grade="高一",
        course_code="C4",
        course_name="课程4",
        quota=5,
        is_deleted=False,
    ))
    db.session.commit()
    plan = preview_course_selection_workbook(
        workbook_bytes(expected, ("C2", None)), expectation=expected,
    )

    async def forbidden(*_args, **_kwargs):
        raise AssertionError("stale course catalog must block the whole batch")

    monkeypatch.setattr(service, "apply_selection_mutations", forbidden)
    with pytest.raises(CourseSelectionWorkbookServiceError) as caught:
        asyncio.run(apply_course_selection_workbook_plan(
            db, plan, expected_plan_id=plan.plan_id, source_expectation=expected,
        ))

    assert caught.value.code == "XBK_WORKBOOK_STALE_BASELINE"
    assert [(issue.code, issue.course_code) for issue in caught.value.issues] == [
        ("XBK_WORKBOOK_STALE_COURSE", "C4"),
    ]
    assert active(db) == {"S1": "C1"}
    assert db.rollback_count == 1


def test_new_active_student_with_selection_rejects_before_any_plan_write(db, monkeypatch):
    expected = expectation()
    seed(db, expected, selections=(("S1", "C1"),))
    db.session.add(XbkStudent(
        year=YEAR,
        term=TERM,
        grade="高一",
        class_name="1班",
        student_no="S3",
        name="学生3",
        is_deleted=False,
    ))
    db.session.add(XbkSelection(
        year=YEAR,
        term=TERM,
        grade="高一",
        student_no="S3",
        name="学生3",
        course_code="C2",
        is_deleted=False,
    ))
    db.session.commit()
    plan = preview_course_selection_workbook(
        workbook_bytes(expected, ("C2", None)), expectation=expected,
    )

    async def forbidden(*_args, **_kwargs):
        raise AssertionError("new roster member must block before plan writes")

    monkeypatch.setattr(service, "apply_selection_mutations", forbidden)
    with pytest.raises(CourseSelectionWorkbookServiceError) as caught:
        asyncio.run(apply_course_selection_workbook_plan(
            db, plan, expected_plan_id=plan.plan_id, source_expectation=expected,
        ))

    assert caught.value.code == "XBK_WORKBOOK_STALE_BASELINE"
    assert [(issue.code, issue.student_no) for issue in caught.value.issues] == [
        ("XBK_WORKBOOK_STALE_STUDENT", "S3"),
    ]
    assert active(db) == {"S1": "C1", "S3": "C2"}
    assert db.rollback_count == 1


def test_changed_rows_at_baseline_are_rejected_when_no_change_row_drifted(db, monkeypatch):
    expected = expectation(baselines=("C1", "C3"))
    seed(db, expected, selections=(("S1", "C1"), ("S2", "C2")))
    plan = preview_course_selection_workbook(
        workbook_bytes(expected, ("C2", None)), expectation=expected,
    )

    async def forbidden(*_args, **_kwargs):
        raise AssertionError("stale no-change row must block the whole batch")

    monkeypatch.setattr(service, "apply_selection_mutations", forbidden)
    with pytest.raises(CourseSelectionWorkbookServiceError) as caught:
        asyncio.run(apply_course_selection_workbook_plan(
            db, plan, expected_plan_id=plan.plan_id, source_expectation=expected,
        ))

    assert caught.value.code == "XBK_WORKBOOK_STALE_BASELINE"
    assert len(caught.value.issues) == 1
    assert caught.value.issues[0].code == "XBK_WORKBOOK_STALE_SELECTION"
    assert caught.value.issues[0].student_no == "S2"
    assert caught.value.issues[0].details == {
        "baseline_course": "C3",
        "current_course": "C2",
    }
    assert active(db) == {"S1": "C1", "S2": "C2"}
    assert db.rollback_count == 1


def test_already_applied_rejects_target_state_that_is_over_quota(db):
    expected = expectation(baselines=("C1", "C3"), quota=1)
    seed(db, expected, selections=(("S1", "C2"), ("S2", "C2")))
    unsigned = service.CourseSelectionWorkbookPlan(
        format=service.WORKBOOK_FORMAT,
        version=service.WORKBOOK_VERSION,
        year=YEAR,
        term=TERM,
        baseline_id=expected.baseline_id,
        rows=(
            service.WorkbookPlannedSelection(
                sheet="高一1班", row=2, grade="高一", class_name="1班",
                student_no="S1", name="学生1", baseline_course="C1",
                submitted_course="C2", target_course="C2", action="select",
            ),
            service.WorkbookPlannedSelection(
                sheet="高一1班", row=3, grade="高一", class_name="1班",
                student_no="S2", name="学生2", baseline_course="C3",
                submitted_course="C2", target_course="C2", action="select",
            ),
        ),
        plan_id="",
    )
    plan = replace(unsigned, plan_id=service._plan_digest(service._plan_payload(unsigned)))

    with pytest.raises(CourseSelectionWorkbookServiceError) as caught:
        asyncio.run(apply_course_selection_workbook_plan(
            db, plan, expected_plan_id=plan.plan_id, source_expectation=expected,
        ))

    assert caught.value.code == "XBK_WORKBOOK_APPLY_CONFLICT"
    assert len(caught.value.issues) == 1
    assert caught.value.issues[0].code == "XBK_WORKBOOK_CAPACITY_CONFLICT"
    assert caught.value.issues[0].course_code == "C2"
    assert caught.value.issues[0].details == {
        "grade": "高一",
        "class_name": "1班",
        "used": 2,
        "quota": 1,
    }
    assert active(db) == {"S1": "C2", "S2": "C2"}
    assert db.rollback_count == 1


def test_changed_rows_at_target_do_not_mask_drifted_no_change_row(db):
    expected = expectation(baselines=("C1", "C3"))
    seed(db, expected, selections=(("S1", "C2"), ("S2", "C2")))
    plan = preview_course_selection_workbook(
        workbook_bytes(expected, ("C2", None)), expectation=expected,
    )

    with pytest.raises(CourseSelectionWorkbookServiceError) as caught:
        asyncio.run(apply_course_selection_workbook_plan(
            db, plan, expected_plan_id=plan.plan_id, source_expectation=expected,
        ))

    assert caught.value.code == "XBK_WORKBOOK_STALE_BASELINE"
    assert len(caught.value.issues) == 1
    assert caught.value.issues[0].code == "XBK_WORKBOOK_STALE_SELECTION"
    assert caught.value.issues[0].student_no == "S2"
    assert active(db) == {"S1": "C2", "S2": "C2"}
    assert db.rollback_count == 1
