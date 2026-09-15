"""Pure synthetic XLSX coverage for the versioned all-class XBK workbook."""

import asyncio
from io import BytesIO
import stat
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import pytest
from openpyxl import load_workbook

from app.models import XbkCourse, XbkSelection, XbkStudent
from app.services.xbk.course_selection_workbook import (
    CATALOG_SHEET,
    INSTRUCTION,
    METADATA_SHEET,
    WorkbookArchiveError,
    WorkbookCourse,
    WorkbookExpectation,
    WorkbookStudent,
    WorkbookValidationError,
    compute_baseline_id,
    parse_course_selection_workbook,
    validate_xlsx_archive_budget,
)
from app.services.xbk.exports.course_selection import (
    CourseSelectionExportConflict,
    build_student_course_selection_xlsx,
)


class Rows:
    def __init__(self, rows):
        self.rows = rows

    def all(self):
        return self.rows

    def scalars(self):
        return self


class CaptureDb:
    def __init__(self, *rows):
        self.results = iter(rows)
        self.statements = []

    async def execute(self, statement):
        self.statements.append(statement)
        return Rows(next(self.results))


def course(code: str, quota: int, grade: str = "高一") -> XbkCourse:
    return XbkCourse(
        year="2026-2027",
        term="上学期",
        grade=grade,
        course_code=code,
        course_name=f"课程{code}",
        teacher="教师",
        quota=quota,
        location="教室",
        is_deleted=False,
    )


def student(no: str, grade: str, class_name: str, name: str | None = None) -> XbkStudent:
    return XbkStudent(
        year="2026-2027",
        term="上学期",
        grade=grade,
        class_name=class_name,
        student_no=no,
        name=name or f"学生{no}",
        is_deleted=False,
    )


def selection(no: str, code: str) -> XbkSelection:
    return XbkSelection(
        year="2026-2027",
        term="上学期",
        student_no=no,
        course_code=code,
        name=f"学生{no}",
        is_deleted=False,
    )


def build(courses, students, selections=(), *, grade="高一", class_name="1班") -> BytesIO:
    db = CaptureDb(list(courses), list(students), list(selections))
    return asyncio.run(build_student_course_selection_xlsx(
        db,
        "2026-2027",
        "上学期",
        class_name,
        None,
        None,
        grade=grade,
    ))


def save_workbook(workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def zip_bytes(entries: list[tuple[str | ZipInfo, bytes]]) -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name, payload in entries:
            archive.writestr(name, payload)
    return output.getvalue()


def mark_first_zip_entry_encrypted(content: bytes) -> bytes:
    """Set the encryption flag in both headers without needing a ZIP encryptor."""

    data = bytearray(content)
    local_header = data.find(b"PK\x03\x04")
    central_header = data.find(b"PK\x01\x02")
    assert local_header >= 0 and central_header >= 0
    for offset in (local_header + 6, central_header + 8):
        flags = int.from_bytes(data[offset : offset + 2], "little") | 0x1
        data[offset : offset + 2] = flags.to_bytes(2, "little")
    return bytes(data)


def test_xlsx_archive_budget_rejects_high_compression_ratio(monkeypatch):
    monkeypatch.setattr(
        "app.services.xbk.course_selection_workbook.MAX_XLSX_COMPRESSION_RATIO", 2
    )
    content = zip_bytes([("xl/worksheets/sheet1.xml", b"A" * 4096)])

    with pytest.raises(WorkbookArchiveError) as caught:
        validate_xlsx_archive_budget(content)

    assert caught.value.status_code == 413
    assert caught.value.code == "XBK_WORKBOOK_ARCHIVE_TOO_LARGE"
    assert "压缩比" in caught.value.message


def test_xlsx_archive_budget_rejects_small_archive_with_large_member(monkeypatch):
    monkeypatch.setattr(
        "app.services.xbk.course_selection_workbook.MAX_XLSX_MEMBER_UNCOMPRESSED_BYTES", 128
    )
    monkeypatch.setattr(
        "app.services.xbk.course_selection_workbook.MAX_XLSX_COMPRESSION_RATIO", 10000
    )
    content = zip_bytes([("xl/sharedStrings.xml", b"B" * 1024)])
    assert len(content) < 512

    with pytest.raises(WorkbookArchiveError) as caught:
        validate_xlsx_archive_budget(content)

    assert caught.value.status_code == 413
    assert "单个条目" in caught.value.message


def test_xlsx_archive_budget_rejects_too_many_entries(monkeypatch):
    monkeypatch.setattr(
        "app.services.xbk.course_selection_workbook.MAX_XLSX_ARCHIVE_ENTRIES", 2
    )
    content = zip_bytes([(f"xl/item-{index}.xml", b"x") for index in range(3)])

    with pytest.raises(WorkbookArchiveError) as caught:
        validate_xlsx_archive_budget(content)

    assert caught.value.status_code == 413
    assert "条目" in caught.value.message


def test_xlsx_archive_budget_rejects_total_uncompressed_size(monkeypatch):
    monkeypatch.setattr(
        "app.services.xbk.course_selection_workbook.MAX_XLSX_TOTAL_UNCOMPRESSED_BYTES", 100
    )
    monkeypatch.setattr(
        "app.services.xbk.course_selection_workbook.MAX_XLSX_COMPRESSION_RATIO", 10000
    )
    content = zip_bytes([("xl/a.xml", b"a" * 60), ("xl/b.xml", b"b" * 60)])

    with pytest.raises(WorkbookArchiveError) as caught:
        validate_xlsx_archive_budget(content)

    assert caught.value.status_code == 413
    assert "总大小" in caught.value.message


def test_xlsx_archive_budget_rejects_too_many_worksheets(monkeypatch):
    monkeypatch.setattr(
        "app.services.xbk.course_selection_workbook.MAX_XLSX_WORKSHEETS", 1
    )
    content = zip_bytes(
        [
            ("xl/worksheets/sheet1.xml", b"<x/>"),
            ("xl/worksheets/sheet2.xml", b"<x/>"),
        ]
    )

    with pytest.raises(WorkbookArchiveError) as caught:
        validate_xlsx_archive_budget(content)

    assert caught.value.status_code == 413
    assert "工作表" in caught.value.message


def test_xlsx_archive_budget_rejects_zip_slip_and_symlink_members():
    with pytest.raises(WorkbookArchiveError) as slip:
        validate_xlsx_archive_budget(zip_bytes([("../escape.xml", b"x")]))
    assert slip.value.status_code == 422

    link = ZipInfo("xl/link.xml")
    link.create_system = 3
    link.external_attr = (stat.S_IFLNK | 0o777) << 16
    with pytest.raises(WorkbookArchiveError) as symlink:
        validate_xlsx_archive_budget(zip_bytes([(link, b"target")]))
    assert symlink.value.status_code == 422


def test_xlsx_archive_budget_rejects_encrypted_members():
    encrypted = mark_first_zip_entry_encrypted(
        zip_bytes([("xl/worksheets/sheet1.xml", b"<x/>")])
    )

    with pytest.raises(WorkbookArchiveError) as caught:
        validate_xlsx_archive_budget(encrypted)

    assert caught.value.status_code == 422
    assert caught.value.code == "XBK_WORKBOOK_ARCHIVE_INVALID"
    assert "加密" in caught.value.message


def test_xlsx_archive_budget_accepts_normal_export():
    output = build(
        [course("C1", 2)],
        [student("0001", "高一", "1班")],
    )

    validate_xlsx_archive_budget(output.getvalue())


def test_one_workbook_uses_every_actual_class_and_strict_simple_layout():
    output = build(
        [course("C1", 2), course("C2", 1), course("C3", 3, grade="高二")],
        [
            student("0001", "高一", "1班"),
            student("0002", "高一", "1班"),
            student("0003", "高一", "5班"),
            student("0004", "高二", "1班"),
        ],
        [selection("0001", "C1")],
    )
    workbook = load_workbook(output, data_only=False)

    visible = [sheet.title for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
    assert visible == [CATALOG_SHEET, "高一1班", "高一5班", "高二1班"]
    assert workbook[METADATA_SHEET].sheet_state == "veryHidden"

    student_numbers = []
    for sheet_name in visible[1:]:
        sheet = workbook[sheet_name]
        assert sheet["F1"].value == INSTRUCTION
        assert [sheet.cell(2, col).value for col in range(6, 11)] == [
            "课程代码", "课程名称", "本班限额", "已选", "剩余",
        ]
        assert "顺序" not in [cell.value for row in sheet.iter_rows() for cell in row]
        assert sheet.protection.sheet
        student_numbers.extend(
            sheet.cell(row, 2).value
            for row in range(2, sheet.max_row + 1)
            if sheet.cell(row, 2).value
        )
    assert sorted(student_numbers) == ["0001", "0002", "0003", "0004"]

    class_one = workbook["高一1班"]
    summary_by_code = {
        class_one.cell(row, 6).value: (
            class_one.cell(row, 9).value,
            class_one.cell(row, 10).value,
            class_one.cell(row, 9).data_type,
            class_one.cell(row, 10).data_type,
        )
        for row in range(3, 3 + 3)
    }
    assert summary_by_code["C1"] == (1, 1, "n", "n")
    assert summary_by_code["C2"] == (0, 1, "n", "n")
    assert summary_by_code["C3"] == (0, 3, "n", "n")
    assert all(
        cell.data_type != "f"
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )
    assert len(class_one.data_validations.dataValidation) == 2
    first_validation = class_one.data_validations.dataValidation[0]
    assert first_validation.allowBlank
    assert 'D2=""' in first_validation.formula1
    assert '="未选"' in first_validation.formula1

    parsed = parse_course_selection_workbook(output)
    assert len(parsed.rows) == 4
    assert {row.action for row in parsed.rows} == {"no_change"}
    assert parsed.year == "2026-2027" and parsed.term == "上学期"


def test_safe_sheet_name_collisions_get_unique_manifest_mapping():
    output = build(
        [course("C1", 2)],
        [
            student("0001", "自定义年级", "超长班级名称ABCDEFGHIJKLMNO/同名"),
            student("0002", "自定义年级", "超长班级名称ABCDEFGHIJKLMNO\\同名"),
        ],
    )
    workbook = load_workbook(output)
    class_sheets = [name for name in workbook.sheetnames if name not in {CATALOG_SHEET, METADATA_SHEET}]
    assert len(class_sheets) == 2
    assert len(set(class_sheets)) == 2
    assert all(len(name) <= 31 for name in class_sheets)
    assert len(parse_course_selection_workbook(output).rows) == 2


def test_export_blocks_existing_multi_selection_instead_of_choosing_first():
    db = CaptureDb(
        [course("C1", 2), course("C2", 2)],
        [student("0001", "高一", "1班")],
        [selection("0001", "C1"), selection("0001", "C2")],
    )
    with pytest.raises(CourseSelectionExportConflict) as exc:
        asyncio.run(build_student_course_selection_xlsx(
            db, "2026-2027", "上学期", None, None, None,
        ))
    assert exc.value.conflicts == {"0001": ["C1", "C2"]}


def test_blank_is_no_change_and_explicit_unselected_is_a_separate_action():
    output = build(
        [course("C1", 2), course("C2", 2)],
        [student("0001", "高一", "1班"), student("0002", "高一", "1班")],
        [selection("0001", "C1")],
    )
    workbook = load_workbook(output)
    sheet = workbook["高一1班"]
    sheet["D2"] = None
    sheet["D3"] = "未选"

    parsed = parse_course_selection_workbook(save_workbook(workbook))
    rows = {row.student_no: row for row in parsed.rows}
    assert rows["0001"].baseline_course == "C1"
    assert rows["0001"].submitted_course is None
    assert rows["0001"].action == "no_change"
    assert rows["0002"].action == "set_unselected"


def test_legal_row_reordering_keeps_identity_and_complete_roster():
    output = build(
        [course("C1", 2)],
        [
            student("0001", "高一", "1班", "甲"),
            student("0002", "高一", "1班", "乙"),
        ],
    )
    workbook = load_workbook(output)
    sheet = workbook["高一1班"]
    first = [sheet.cell(2, column).value for column in range(1, 5)]
    second = [sheet.cell(3, column).value for column in range(1, 5)]
    for column, value in enumerate(second, 1):
        sheet.cell(2, column).value = value
    for column, value in enumerate(first, 1):
        sheet.cell(3, column).value = value

    parsed = parse_course_selection_workbook(save_workbook(workbook))
    assert {row.student_no: row.row for row in parsed.rows} == {"0002": 2, "0001": 3}


def test_capacity_conflicts_include_sheet_row_student_course_and_same_class_advice():
    output = build(
        [course("C1", 1), course("C2", 2)],
        [
            student("0001", "高一", "1班"),
            student("0002", "高一", "1班"),
            student("0003", "高一", "2班"),
            student("0004", "高一", "2班"),
        ],
        [selection("0003", "C2"), selection("0004", "C2")],
    )
    workbook = load_workbook(output)
    workbook["高一1班"]["D2"] = "C1"
    workbook["高一1班"]["D3"] = "C1"

    with pytest.raises(WorkbookValidationError) as exc:
        parse_course_selection_workbook(save_workbook(workbook))
    capacity = [issue for issue in exc.value.issues if "超过限额" in issue.reason]
    assert {(issue.sheet, issue.row, issue.student_no, issue.course_code) for issue in capacity} == {
        ("高一1班", 2, "0001", "C1"),
        ("高一1班", 3, "0002", "C1"),
    }
    assert all(issue.suggestions[0].course_code == "C2" for issue in capacity)
    # C2 is full in another class, but this class still has its own two places.
    assert all(issue.suggestions[0].remaining == 2 for issue in capacity)


def test_missing_sheet_unknown_copy_and_identity_change_are_all_reported():
    output = build(
        [course("C1", 2), course("C2", 2)],
        [student("0001", "高一", "1班"), student("0002", "高一", "2班")],
    )
    workbook = load_workbook(output)
    del workbook["高一2班"]
    copy = workbook.copy_worksheet(workbook["高一1班"])
    copy.title = "复制班级"
    workbook["高一1班"]["A2"] = "9班"
    workbook["高一1班"]["D2"] = "不存在"

    with pytest.raises(WorkbookValidationError) as exc:
        parse_course_selection_workbook(save_workbook(workbook))
    payload = exc.value.to_dict()
    assert all({"sheet", "row", "student_no", "course_code", "reason", "suggestions"} <= set(item) for item in payload["issues"])
    reasons = {issue.reason for issue in exc.value.issues}
    assert "缺少导出的班级工作表" in reasons
    assert "包含未知或复制的工作表" in reasons
    assert "学生班级或姓名与导出名册不一致" in reasons
    assert "课程代码不在导出目录中" in reasons


def test_catalog_protection_period_and_server_baseline_are_not_trusted_from_cells():
    output = build(
        [course("C1", 2)],
        [student("0001", "高一", "1班")],
    )
    original = parse_course_selection_workbook(output)
    expectation = WorkbookExpectation(
        year=original.year,
        term=original.term,
        baseline_id=original.baseline_id,
        students=(),
        courses=(),
    )
    workbook = load_workbook(output)
    workbook[CATALOG_SHEET].protection.sheet = False
    workbook[CATALOG_SHEET]["B3"] = "被修改"
    workbook[METADATA_SHEET]["B4"] = "下学期"

    with pytest.raises(WorkbookValidationError) as exc:
        parse_course_selection_workbook(save_workbook(workbook), expected=expectation)
    reasons = {issue.reason for issue in exc.value.issues}
    assert "课程目录保护已被移除" in reasons
    assert "课程目录与导出基线不一致" in reasons
    assert any("学年或学期" in reason or "导出基线" in reason for reason in reasons)


def test_valid_server_expectation_accepts_the_complete_export_baseline():
    output = build(
        [course("C1", 2)],
        [student("0001", "高一", "1班", "甲")],
        [selection("0001", "C1")],
    )
    students = (
        WorkbookStudent("高一1班", "高一", "1班", "0001", "甲", "C1"),
    )
    courses = (WorkbookCourse("C1", "课程C1", 2),)
    expectation = WorkbookExpectation(
        year="2026-2027",
        term="上学期",
        baseline_id=compute_baseline_id("2026-2027", "上学期", students, courses),
        students=students,
        courses=courses,
    )

    parsed = parse_course_selection_workbook(output, expected=expectation)

    assert parsed.baseline_id == expectation.baseline_id
    assert len(parsed.rows) == 1


def test_empty_period_keeps_a_visible_catalog_and_parses_as_an_empty_complete_set():
    output = build([], [])
    workbook = load_workbook(output)

    assert workbook[CATALOG_SHEET].sheet_state == "visible"
    assert workbook[METADATA_SHEET].sheet_state == "veryHidden"
    assert parse_course_selection_workbook(output).rows == ()
