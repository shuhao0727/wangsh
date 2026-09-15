"""Real PostgreSQL concurrency/cancellation proof for the R3 workbook confirm path.

Opt-in only: ``TEST_DATABASE_URL`` must point at a dedicated loopback PostgreSQL
*test* database on a port other than 5432/5433. Every test creates a unique
schema, uses synthetic rows, and drops that schema in cleanup.
"""

from __future__ import annotations

import asyncio
from io import BytesIO
import os
import re
import unittest
import uuid

import pytest
from fastapi import HTTPException
from openpyxl import Workbook
from sqlalchemy import func, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.api.endpoints.xbk.selections import create_selection
from app.models import XbkCourse, XbkSelection, XbkStudent
from app.schemas.xbk import XbkSelectionUpsert
from app.services.xbk.course_selection_workbook import (
    CATALOG_SHEET,
    INSTRUCTION,
    METADATA_SHEET,
    WorkbookCourse,
    WorkbookExpectation,
    WorkbookStudent,
    compute_baseline_id,
    write_metadata_sheet,
)
from app.services.xbk.course_selection_workbook_apply import (
    CourseSelectionWorkbookServiceError,
    apply_course_selection_workbook_plan,
    preview_course_selection_workbook,
)

YEAR, TERM = "2037-2038", "上学期"


def dedicated_url() -> str:
    value = os.environ.get("TEST_DATABASE_URL", "").strip()
    if not value:
        pytest.skip("requires explicitly allocated PostgreSQL test database")
    parsed = make_url(value)
    database = parsed.database or ""
    host = (parsed.host or "").lower()
    if parsed.drivername != "postgresql+asyncpg":
        pytest.fail("refusing non-asyncpg database before connecting")
    if host not in {"127.0.0.1", "localhost", "::1"}:
        pytest.fail("refusing non-loopback PostgreSQL before connecting")
    if parsed.port in {5432, 5433}:
        pytest.fail("refusing protected PostgreSQL ports 5432/5433 before connecting")
    if not re.search(r"(?:^|[_-])(?:test|testing|ci)(?:$|[_-])", database, re.I):
        pytest.fail("refusing non-test database before connecting")
    return value


def expectation(
    *, quota: int = 1, baselines: tuple[str | None, str | None] = (None, None)
) -> WorkbookExpectation:
    students = (
        WorkbookStudent("高一1班", "高一", "1班", "S1", "学生1", baselines[0]),
        WorkbookStudent("高一1班", "高一", "1班", "S2", "学生2", baselines[1]),
    )
    courses = (
        WorkbookCourse("C1", "课程1", quota),
        WorkbookCourse("C2", "课程2", 5),
    )
    return WorkbookExpectation(
        year=YEAR,
        term=TERM,
        baseline_id=compute_baseline_id(YEAR, TERM, students, courses),
        students=students,
        courses=courses,
    )


def workbook_bytes(expected: WorkbookExpectation, submitted: tuple[str | None, str | None]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    catalog = wb.create_sheet(CATALOG_SHEET)
    catalog.append(["课程目录"])
    catalog.append(["课程代码", "课程名称", "课程负责人", "各班限报人数", "上课地点"])
    for course in expected.courses:
        catalog.append([course.course_code, course.course_name, None, course.quota, None])
    catalog.protection.sheet = True

    ws = wb.create_sheet("高一1班")
    ws.append(["班级", "学号", "姓名", "课程代码", None, INSTRUCTION])
    ws.cell(2, 6, "课程代码")
    ws.cell(2, 7, "课程名称")
    ws.cell(2, 8, "本班限额")
    ws.cell(2, 9, "已选")
    ws.cell(2, 10, "剩余")
    ws.protection.sheet = True
    for row_no, (student, chosen) in enumerate(zip(expected.students, submitted), start=2):
        ws.cell(row_no, 1, student.class_name)
        ws.cell(row_no, 2, student.student_no)
        ws.cell(row_no, 3, student.name)
        ws.cell(row_no, 4, chosen)

    metadata = wb.create_sheet(METADATA_SHEET)
    write_metadata_sheet(
        metadata,
        year=expected.year,
        term=expected.term,
        students=expected.students,
        courses=expected.courses,
    )
    metadata.sheet_state = "veryHidden"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def plan_for(expected: WorkbookExpectation, submitted: tuple[str | None, str | None]):
    return preview_course_selection_workbook(workbook_bytes(expected, submitted), expectation=expected)


class WorkbookPostgresConcurrency(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self) -> None:
        self.schema = "test_r3_workbook_" + uuid.uuid4().hex
        url = dedicated_url()
        self.admin = create_async_engine(url)
        self.engine = create_async_engine(
            url,
            connect_args={
                "server_settings": {
                    "search_path": self.schema,
                    "application_name": self.schema,
                    "statement_timeout": "15000",
                    "lock_timeout": "12000",
                }
            },
        )
        self.Session = async_sessionmaker(self.engine, expire_on_commit=False)
        async with self.admin.begin() as conn:
            await conn.execute(text(f'CREATE SCHEMA "{self.schema}"'))
        try:
            async with self.engine.begin() as conn:
                for model in (XbkStudent, XbkCourse, XbkSelection):
                    await conn.run_sync(model.__table__.create)
        except BaseException:
            await self._cleanup_resources()
            raise

    async def asyncTearDown(self) -> None:
        await self._cleanup_resources()

    async def _cleanup_resources(self) -> None:
        engine = getattr(self, "engine", None)
        admin = getattr(self, "admin", None)
        schema = getattr(self, "schema", None)
        if engine is not None:
            await engine.dispose()
        if admin is not None and schema is not None:
            async with admin.begin() as conn:
                await conn.execute(text(
                    "SELECT pg_terminate_backend(pid) FROM pg_stat_activity "
                    "WHERE application_name=:name AND pid <> pg_backend_pid()"
                ), {"name": schema})
                await conn.execute(text(f'DROP SCHEMA IF EXISTS "{schema}" CASCADE'))
            await admin.dispose()

    async def seed(self, expected: WorkbookExpectation, *, selections=()) -> None:
        async with self.Session() as db:
            db.add_all([
                XbkStudent(
                    year=YEAR,
                    term=TERM,
                    grade=student.grade,
                    class_name=student.class_name,
                    student_no=student.student_no,
                    name=student.name,
                    is_deleted=False,
                )
                for student in expected.students
            ])
            db.add_all([
                XbkCourse(
                    year=YEAR,
                    term=TERM,
                    grade="高一",
                    course_code=course.course_code,
                    course_name=course.course_name,
                    quota=course.quota,
                    is_deleted=False,
                )
                for course in expected.courses
            ])
            db.add_all([
                XbkSelection(
                    year=YEAR,
                    term=TERM,
                    grade="高一",
                    student_no=student_no,
                    name=student_no,
                    course_code=course_code,
                    is_deleted=False,
                )
                for student_no, course_code in selections
            ])
            await db.commit()

    async def active(self) -> dict[str, str]:
        async with self.Session() as db:
            rows = (await db.execute(
                select(XbkSelection).where(XbkSelection.is_deleted.is_(False))
            )).scalars().all()
            return {row.student_no: row.course_code for row in rows}

    async def wait_for_backend(self, *, wait_event_type: str | None = None, query: str | None = None) -> int:
        for _ in range(400):
            async with self.admin.connect() as conn:
                rows = (await conn.execute(text(
                    "SELECT pid, wait_event_type, wait_event, query FROM pg_stat_activity "
                    "WHERE application_name=:name AND state <> 'idle'"
                ), {"name": self.schema})).all()
            for pid, event_type, _event, current_query in rows:
                if wait_event_type and event_type != wait_event_type:
                    continue
                if query and query.upper() not in (current_query or "").upper():
                    continue
                return int(pid)
            await asyncio.sleep(0.02)
        self.fail(f"database wait not observed: type={wait_event_type!r}, query={query!r}")

    async def assert_backend_released(self, pid: int) -> None:
        for _ in range(300):
            async with self.admin.connect() as conn:
                exists = bool((await conn.execute(
                    text("SELECT count(*) FROM pg_stat_activity WHERE pid=:pid"), {"pid": pid}
                )).scalar_one())
            if not exists:
                return
            await asyncio.sleep(0.02)
        self.fail(f"PostgreSQL backend {pid} remained after cancellation")

    async def confirm(self, db: AsyncSession, plan, expected: WorkbookExpectation, *, commit=True):
        result = await apply_course_selection_workbook_plan(
            db,
            plan,
            expected_plan_id=plan.plan_id,
            source_expectation=expected,
        )
        if commit:
            await db.commit()
        return result

    async def install_insert_sleep(self) -> None:
        async with self.engine.begin() as conn:
            await conn.execute(text("""
                CREATE FUNCTION r3_workbook_insert_gate() RETURNS trigger LANGUAGE plpgsql AS $$
                BEGIN
                    IF NEW.student_no = 'S1' THEN PERFORM pg_sleep(30); END IF;
                    RETURN NEW;
                END $$
            """))
            await conn.execute(text("""
                CREATE TRIGGER r3_workbook_insert_gate
                BEFORE INSERT ON xbk_selections
                FOR EACH ROW EXECUTE FUNCTION r3_workbook_insert_gate()
            """))

    async def drop_insert_sleep(self) -> None:
        async with self.engine.begin() as conn:
            await conn.execute(text("DROP TRIGGER IF EXISTS r3_workbook_insert_gate ON xbk_selections"))
            await conn.execute(text("DROP FUNCTION IF EXISTS r3_workbook_insert_gate()"))

    async def test_two_workbooks_serialize_and_second_rechecks_full_baseline(self) -> None:
        expected = expectation(quota=1)
        await self.seed(expected)
        first_plan = plan_for(expected, ("C1", None))
        second_plan = plan_for(expected, (None, "C1"))
        applied = asyncio.Event()
        release = asyncio.Event()

        async def first():
            async with self.Session() as db:
                result = await self.confirm(db, first_plan, expected, commit=False)
                applied.set()
                await asyncio.wait_for(release.wait(), 8)
                await db.commit()
                return result

        async def second():
            async with self.Session() as db:
                return await self.confirm(db, second_plan, expected)

        task1 = asyncio.create_task(first())
        await asyncio.wait_for(applied.wait(), 6)
        task2 = asyncio.create_task(second())
        await self.wait_for_backend(wait_event_type="Lock")
        release.set()
        results = await asyncio.wait_for(asyncio.gather(task1, task2, return_exceptions=True), 12)

        self.assertEqual(results[0].status, "applied")
        self.assertIsInstance(results[1], CourseSelectionWorkbookServiceError)
        self.assertEqual(results[1].code, "XBK_WORKBOOK_STALE_BASELINE")
        self.assertEqual(await self.active(), {"S1": "C1"})

    async def test_same_plan_concurrent_confirms_apply_once_then_report_already_applied(self) -> None:
        expected = expectation(quota=1)
        await self.seed(expected)
        plan = plan_for(expected, ("C1", None))
        applied = asyncio.Event()
        release = asyncio.Event()

        async def first():
            async with self.Session() as db:
                result = await self.confirm(db, plan, expected, commit=False)
                applied.set()
                await asyncio.wait_for(release.wait(), 8)
                await db.commit()
                return result

        async def second():
            async with self.Session() as db:
                return await self.confirm(db, plan, expected)

        task1 = asyncio.create_task(first())
        await asyncio.wait_for(applied.wait(), 6)
        task2 = asyncio.create_task(second())
        await self.wait_for_backend(wait_event_type="Lock")
        release.set()
        first_result, second_result = await asyncio.wait_for(
            asyncio.gather(task1, task2),
            12,
        )

        self.assertEqual(first_result.status, "applied")
        self.assertEqual(first_result.changed, 1)
        self.assertEqual(first_result.inserted, 1)
        self.assertEqual(first_result.updated, 0)
        self.assertEqual(second_result.status, "already_applied")
        self.assertEqual(second_result.changed, 0)
        self.assertEqual(second_result.inserted, 0)
        self.assertEqual(second_result.updated, 0)

        async with self.Session() as db:
            active_rows = (await db.execute(
                select(XbkSelection).where(
                    XbkSelection.year == YEAR,
                    XbkSelection.term == TERM,
                    XbkSelection.is_deleted.is_(False),
                )
            )).scalars().all()
            course = (await db.execute(
                select(XbkCourse).where(
                    XbkCourse.year == YEAR,
                    XbkCourse.term == TERM,
                    XbkCourse.course_code == "C1",
                    XbkCourse.is_deleted.is_(False),
                )
            )).scalar_one()

        self.assertEqual(
            [(row.student_no, row.course_code) for row in active_rows],
            [("S1", "C1")],
        )
        self.assertEqual(len({row.student_no for row in active_rows}), len(active_rows))
        self.assertLessEqual(
            sum(row.course_code == "C1" for row in active_rows),
            int(course.quota or 0),
        )

    async def test_workbook_and_single_selection_compete_for_last_seat(self) -> None:
        expected = expectation(quota=1)
        await self.seed(expected)
        plan = plan_for(expected, ("C1", None))
        applied = asyncio.Event()
        release = asyncio.Event()

        async def workbook():
            async with self.Session() as db:
                result = await self.confirm(db, plan, expected, commit=False)
                applied.set()
                await asyncio.wait_for(release.wait(), 8)
                await db.commit()
                return result

        async def ordinary():
            async with self.Session() as db:
                payload = XbkSelectionUpsert(
                    year=YEAR,
                    term=TERM,
                    grade="高一",
                    student_no="S2",
                    name="学生2",
                    course_code="C1",
                )
                return await create_selection(payload, db, {})

        task1 = asyncio.create_task(workbook())
        await asyncio.wait_for(applied.wait(), 6)
        task2 = asyncio.create_task(ordinary())
        await self.wait_for_backend(wait_event_type="Lock")
        release.set()
        results = await asyncio.wait_for(asyncio.gather(task1, task2, return_exceptions=True), 12)

        self.assertEqual(results[0].status, "applied")
        self.assertIsInstance(results[1], HTTPException)
        self.assertEqual(results[1].status_code, 409)
        self.assertEqual(await self.active(), {"S1": "C1"})

    async def test_cancel_while_waiting_for_database_lock_rolls_back_and_releases_backend(self) -> None:
        expected = expectation(quota=2)
        await self.seed(expected)
        plan = plan_for(expected, ("C1", None))
        blocker = await self.engine.connect()
        db = self.Session()
        try:
            await blocker.execute(text("SELECT id FROM xbk_students WHERE student_no='S1' FOR UPDATE"))
            pid = int((await db.execute(text("SELECT pg_backend_pid()"))).scalar_one())
            task = asyncio.create_task(self.confirm(db, plan, expected))
            await self.wait_for_backend(wait_event_type="Lock")
            task.cancel()
            with self.assertRaises(asyncio.CancelledError):
                await task
            self.assertFalse(db.in_transaction())
            await blocker.rollback()
            await db.close()
            await self.assert_backend_released(pid)
            self.assertEqual(await self.active(), {})
        finally:
            if db.in_transaction():
                await db.rollback()
            await db.close()
            await blocker.rollback()
            await blocker.close()

    async def test_cancel_during_real_insert_rolls_back_partial_write(self) -> None:
        expected = expectation(quota=2)
        await self.seed(expected)
        await self.install_insert_sleep()
        plan = plan_for(expected, ("C1", None))
        async with self.Session() as db:
            pid = int((await db.execute(text("SELECT pg_backend_pid()"))).scalar_one())
            task = asyncio.create_task(self.confirm(db, plan, expected))
            await self.wait_for_backend(query="INSERT")
            task.cancel()
            with self.assertRaises(asyncio.CancelledError):
                await task
            self.assertFalse(db.in_transaction())
        await self.assert_backend_released(pid)
        self.assertEqual(await self.active(), {})
        await self.drop_insert_sleep()
        async with self.Session() as retry:
            result = await self.confirm(retry, plan, expected)
            self.assertEqual(result.status, "applied")
        self.assertEqual(await self.active(), {"S1": "C1"})

    async def test_slow_rollback_is_awaited_after_write_cancellation(self) -> None:
        expected = expectation(quota=2)
        await self.seed(expected)
        await self.install_insert_sleep()
        plan = plan_for(expected, ("C1", None))
        rollback_entered = asyncio.Event()
        release_rollback = asyncio.Event()

        class SlowRollbackSession(AsyncSession):
            async def rollback(inner) -> None:
                rollback_entered.set()
                await asyncio.wait_for(release_rollback.wait(), 8)
                await super().rollback()

        db = SlowRollbackSession(self.engine, expire_on_commit=False)
        try:
            task = asyncio.create_task(self.confirm(db, plan, expected))
            await self.wait_for_backend(query="INSERT")
            task.cancel()
            await asyncio.wait_for(rollback_entered.wait(), 5)
            self.assertFalse(task.done(), "cancel returned before protected rollback completed")
            release_rollback.set()
            with self.assertRaises(asyncio.CancelledError):
                await asyncio.wait_for(task, 8)
            self.assertFalse(db.in_transaction())
            self.assertEqual(await self.active(), {})
        finally:
            release_rollback.set()
            if db.in_transaction():
                await AsyncSession.rollback(db)
            await db.close()

    async def test_failed_rollback_preserves_cancelled_error_and_close_discards_transaction(self) -> None:
        expected = expectation(quota=2)
        await self.seed(expected)
        await self.install_insert_sleep()
        plan = plan_for(expected, ("C1", None))
        rollback_called = asyncio.Event()

        class FailingRollbackSession(AsyncSession):
            async def rollback(inner) -> None:
                rollback_called.set()
                raise RuntimeError("synthetic rollback failure")

        db = FailingRollbackSession(self.engine, expire_on_commit=False)
        try:
            task = asyncio.create_task(self.confirm(db, plan, expected))
            await self.wait_for_backend(query="INSERT")
            task.cancel()
            with self.assertRaises(asyncio.CancelledError):
                await task
            await asyncio.wait_for(rollback_called.wait(), 3)
            self.assertTrue(db.in_transaction(), "failed rollback must leave session unusable for caller disposal")
        finally:
            await db.close()
        self.assertEqual(await self.active(), {})

    async def test_confirmation_session_rollback_does_not_touch_other_business_session(self) -> None:
        expected = expectation(quota=1)
        await self.seed(expected)
        plan = plan_for(expected, ("C1", None))
        async with self.engine.begin() as conn:
            await conn.execute(text(
                "ALTER TABLE xbk_selections ADD CONSTRAINT ck_synthetic_confirm_failure "
                "CHECK (student_no <> 'S1')"
            ))
        business = self.Session()
        try:
            pending = XbkStudent(
                year=YEAR,
                term=TERM,
                grade="高一",
                class_name="2班",
                student_no="BUSINESS-1",
                name="独立业务",
                is_deleted=False,
            )
            business.add(pending)
            self.assertIn(pending, business.sync_session.new)

            async with self.Session() as confirmation:
                with self.assertRaises(IntegrityError):
                    await self.confirm(confirmation, plan, expected)
                self.assertFalse(confirmation.in_transaction())

            self.assertTrue(business.in_transaction())
            self.assertIn(pending, business.sync_session.new)
            await business.commit()
        finally:
            await business.close()

        async with self.Session() as verify:
            count = await verify.scalar(select(func.count()).select_from(XbkStudent).where(
                XbkStudent.student_no == "BUSINESS-1"
            ))
            self.assertEqual(count, 1)
        self.assertEqual(await self.active(), {})
