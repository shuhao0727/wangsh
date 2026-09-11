"""Opt-in HTTP regression against the main agent's throwaway Docker audit app.

Never targets port 8000 or production; only writes the allocated academic year 2032-2033.
Run with XBK_AUDIT_BASE_URL=http://localhost:8009 and XBK_AUDIT_ALLOW_LEGACY_SEED=1.
The scratch adapter alone implements POST /__audit__/xbk/legacy-empty-codes:
accept {year, term, tag}, require the isolated PG and audit admin, seed only the
matching test roster, and return committed empty-code rows from a fresh session.
This fixture API must never be added to the normal application.
The dedicated audit app injects identities; this is not an authentication test.
"""
import os
from io import BytesIO
from uuid import uuid4

import httpx
import pytest
from openpyxl import load_workbook


AUDIT_ACADEMIC_YEAR = "2032-2033"


@pytest.mark.skipif(not os.environ.get("XBK_AUDIT_BASE_URL"), reason="requires explicitly allocated isolated XBK audit app")
def test_nonimport_http_regression_on_disposable_2032_data():
    base = os.environ["XBK_AUDIT_BASE_URL"].rstrip("/")
    assert base in {"http://localhost:8009", "http://127.0.0.1:8009"}, "refuse any non-audit endpoint"
    year, term = AUDIT_ACADEMIC_YEAR, "上学期"
    tag = uuid4().hex[:8]
    api = base + "/api/v1/xbk"
    outcomes = []
    with httpx.Client(timeout=30, headers={"x-audit-role": "admin"}) as client:
        assert client.get(base + "/health").json() == {"status": "isolated-xbk-test"}

        def request(method, path, expected=200, **kwargs):
            response = client.request(method, api + path, **kwargs)
            assert response.status_code == expected, (method, path, response.status_code, response.text[:800])
            outcomes.append((method, path, expected))
            return response

        def student(no, grade="高一", cls="审查A班", semester=term):
            payload = dict(year=year, term=semester, grade=grade, class_name=cls, student_no=no, name=no)
            return payload, request("POST", "/data/students", json=payload).json()

        def course(code, grade="高一"):
            payload = dict(year=year, term=term, grade=grade, course_code=code, course_name=code, quota=3)
            return payload, request("POST", "/data/courses", json=payload).json()

        def selection(no, code, grade="高一"):
            payload = dict(year=year, term=term, grade=grade, student_no=no, name=no, course_code=code)
            return payload, request("POST", "/data/selections", json=payload).json()

        # Dependency injection only replaces identity retrieval. require_admin remains real.
        payload = dict(year=year, term=term, class_name="审查权限班", student_no="forbidden", name="forbidden")
        for role, status in [("student", 403), ("teacher", 403), ("anonymous", 401)]:
            request("POST", "/data/students", status, headers={"x-audit-role": role}, json=payload)
            request("DELETE", "/data", status, headers={"x-audit-role": role}, params=dict(scope="all", year=year, term=term))
            request("GET", "/export/distribution", status, headers={"x-audit-role": role}, params=dict(year=year, term=term))
        enabled = request("GET", "/public-config").json()["enabled"]
        request("GET", "/data/students", 200 if enabled else 403,
                headers={"x-audit-role": "anonymous"}, params={"year": year})

        p1, s1 = student("A" + tag)
        p2, s2 = student("B" + tag)
        p3, s3 = student("C" + tag, "高二", "审查A班")
        student(p1["student_no"], semester="下学期")
        pc1, c1 = course("COURSE_A" + tag)
        pc2, c2 = course("COURSE_B" + tag, "高二")
        pc0, c0 = course("COURSE_ZERO" + tag)
        pq1, q1 = selection(p1["student_no"], pc1["course_code"], grade="高二")  # stale snapshot
        pq2, q2 = selection(p3["student_no"], pc2["course_code"], grade="高二")
        scope = dict(year=year, term=term, grade="高一", class_name="审查A班")
        summary = request("GET", "/analysis/summary", params=scope).json()
        assert summary["students"] >= 2 and summary["selections"] >= 1
        stats = request("GET", "/analysis/course-stats", params=scope).json()["items"]
        assert next(r for r in stats if r["course_code"] == pc1["course_code"])["count"] == 1
        zero_row = next(r for r in stats if r["course_code"] == pc0["course_code"])
        assert zero_row["count"] == 0 and zero_row["grade"] == "高一"
        real_rows = [r for r in stats if r["course_code"] not in {"未选", "休学或其他"}]
        assert summary["courses"] == len(real_rows) == 2

        # Ordinary selection exports use the live roster grade for valid students,
        # so stale grade snapshots cannot disagree with the page filters/statistics.
        for export_scope in ["selections", "course_results"]:
            response = request("GET", "/export", params={**scope, "scope": export_scope})
            values = list(load_workbook(BytesIO(response.content)).active.values)
            records = [dict(zip(values[0], row)) for row in values[1:]]
            # The data export mirrors the list's active roster: a student without
            # any live selection still appears as a virtual "休学或其他" row.
            assert (p2["student_no"], "休学或其他") in {
                (row["学号"], row["课程代码"]) for row in records
            }
            real_rows = [
                (row["学号"], row["年级"]) for row in records
                if row["课程代码"] not in {"未选", "休学或其他"}
            ]
            assert real_rows == [
                (p1["student_no"], "高一"),
            ]

        # The same class name can exist in multiple grades. Class-only filtering
        # must retain each grade's catalog while calculating capacity per grade.
        cross_scope = dict(year=year, term=term, class_name="审查A班")
        cross_summary = request("GET", "/analysis/summary", params=cross_scope).json()
        cross_stats = request("GET", "/analysis/course-stats", params=cross_scope).json()["items"]
        cross_real_rows = {
            r["course_code"]: r
            for r in cross_stats
            if r["course_code"] not in {"未选", "休学或其他"}
        }
        assert cross_summary["students"] == 3
        assert cross_summary["courses"] == len(cross_real_rows) == 3
        assert {(r["grade"], r["class_count"]) for r in cross_real_rows.values()} == {
            ("高一", 1), ("高二", 1),
        }
        assert cross_real_rows[pc2["course_code"]]["count"] == 1
        absent = request("GET", "/analysis/students-without-selection", params=dict(year=year)).json()["items"]
        assert any(r["student_no"] == p1["student_no"] and r["term"] == "下学期" for r in absent)

        # Unique conflicts return controlled 409s and preserve the old data.
        request("PUT", f'/data/students/{s1["id"]}', 409, json=p2)
        request("PUT", f'/data/courses/{c1["id"]}', 409, json=pc2)
        _, extra_selection = selection(p1["student_no"], pc2["course_code"])
        request("PUT", f'/data/selections/{q1["id"]}', 409,
                json={**pq1, "course_code": pc2["course_code"]})
        request("DELETE", f'/data/selections/{extra_selection["id"]}', 204)
        request("DELETE", "/data/selections/0", 404)
        request("DELETE", "/data", 400, params={"scope": "all", "year": year})
        for kind in ["all", "courses"]:
            request("DELETE", "/data", 400, params={"scope": kind, **scope})

        # Empty workbooks and grade-scoped exports open cleanly with openpyxl.
        for kind in ["course-selection", "distribution", "teacher-distribution"]:
            response = request("GET", "/export/" + kind, params=scope)
            workbook = load_workbook(BytesIO(response.content))
            values = [str(cell.value) for sheet in workbook for row in sheet for cell in row if cell.value is not None]
            assert p3["student_no"] not in values
            assert all(pc2["course_code"] not in value for value in values)
            assert "filename*=UTF-8''" in response.headers["content-disposition"]
        for kind in ["distribution", "teacher-distribution"]:
            response = request("GET", "/export/" + kind,
                               params=dict(year=year, term=term, class_name="不存在的审查班"))
            workbook = load_workbook(BytesIO(response.content))
            assert workbook.active["A1"].value == "当前筛选条件下暂无数据"

        # PostgreSQL regex and cast paths: the field allows up to 50 digits.
        long_no = "9" * 42 + str(int(tag, 16)).zfill(10)[-8:]
        student(long_no)
        course(long_no)
        selection(long_no, long_no)
        for path in ["/data/courses", "/data/course-results"]:
            result = request("GET", path, params=dict(year=year, term=term)).json()
            assert result["total"] >= 1
        for kind in ["course-selection", "distribution", "teacher-distribution"]:
            response = request("GET", "/export/" + kind, params=dict(year=year, term=term))
            assert load_workbook(BytesIO(response.content)).sheetnames

        # Soft-delete cascades do not silently revive when parent records are recreated.
        request("DELETE", f'/data/students/{s1["id"]}', 204)
        request("POST", "/data/students", json=p1)
        visible = request("GET", "/data/selections", params={**scope, "search_text": p1["student_no"]}).json()["items"]
        assert len(visible) == 1 and visible[0]["id"] == 0
        request("DELETE", f'/data/courses/{c2["id"]}', 204)
        request("POST", "/data/courses", json=pc2)
        visible = request("GET", "/data/selections", params=dict(year=year, term=term, search_text=p3["student_no"])).json()["items"]
        assert len(visible) == 1 and visible[0]["id"] == 0

        # Only the allocated 2032-2033 data is physically removed. No other academic years are read or written.
        request("DELETE", "/data", params=dict(scope="all", year=year, term=term))
        request("DELETE", "/data", params=dict(scope="all", year=year, term="下学期"))
        result = request("GET", "/data/students", params={"year": year}).json()
        assert result["total"] == 0
    print(f"\n{len(outcomes)} isolated HTTP assertions passed (academic year {year} only)")


@pytest.mark.skipif(not os.environ.get("XBK_AUDIT_BASE_URL"), reason="requires explicitly allocated isolated XBK audit app")
def test_multipart_blank_selection_import_consistent_across_views():
    """Exercise actual XLSX multipart parsing, persistence, analysis and export."""
    from openpyxl import Workbook

    base = os.environ["XBK_AUDIT_BASE_URL"].rstrip("/")
    assert base in {"http://localhost:8009", "http://127.0.0.1:8009"}
    tag = uuid4().hex[:8]
    scope = dict(year=AUDIT_ACADEMIC_YEAR, term="上学期", grade="高一", class_name="空选回归" + tag)
    chosen, empty, code = "CHOSEN_" + tag, "EMPTY_" + tag, "COURSE_" + tag
    api = base + "/api/v1/xbk"
    course_id = None
    blank_course_id = None
    assert os.environ.get("XBK_AUDIT_ALLOW_LEGACY_SEED") == "1", (
        "requires explicit scratch-only legacy PG fixture support; never use the normal app"
    )
    with httpx.Client(timeout=30, headers={"x-audit-role": "admin"}) as client:
        assert client.get(base + "/health").json() == {"status": "isolated-xbk-test"}

        def request(method, path, expected=200, **kwargs):
            response = client.request(method, api + path, **kwargs)
            assert response.status_code == expected, (method, path, response.status_code, response.text[:800])
            return response

        try:
            for no in [chosen, empty]:
                request("POST", "/data/students", json={**scope, "student_no": no, "name": no})
            course_id = request("POST", "/data/courses", json={
                "year": AUDIT_ACADEMIC_YEAR, "term": scope["term"], "grade": "高一",
                "course_code": code, "course_name": code, "quota": 3,
            }).json()["id"]
            book = Workbook()
            sheet = book.active
            sheet.append(["学年", "学期", "年级", "学号", "姓名", "课程代码"])
            sheet.append([AUDIT_ACADEMIC_YEAR, scope["term"], "高二", chosen, chosen, code])
            sheet.append([AUDIT_ACADEMIC_YEAR, scope["term"], "高二", empty, empty, None])
            stream = BytesIO()
            book.save(stream)
            imported = request("POST", "/import", params={
                "scope": "selections", "year": AUDIT_ACADEMIC_YEAR, "term": scope["term"],
                "grade": "高二", "skip_invalid": "false",
            }, files={"file": ("blank-selection.xlsx", stream.getvalue(),
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}).json()
            assert imported["processed"] == 2, imported

            summary = request("GET", "/analysis/summary", params=scope).json()
            assert {key: summary[key] for key in ["students", "selections", "unselected_count", "suspended_count"]} == {
                "students": 2, "selections": 1, "unselected_count": 1, "suspended_count": 0,
            }
            rows = request("GET", "/analysis/students-with-empty-selection", params=scope).json()["items"]
            assert [row["student_no"] for row in rows] == [empty]
            stats = request("GET", "/analysis/course-stats", params=scope).json()["items"]
            assert sum(row["count"] for row in stats if row["course_code"] in {"", "未选"}) == 1
            assert sum(row["count"] for row in stats if row["course_code"] == code) == 1
            assert sum(row["count"] for row in stats) == 2
            exported = request("GET", "/export", params={**scope, "scope": "unselected", "format": "xlsx"})
            rows = list(load_workbook(BytesIO(exported.content)).active.values)
            assert len(rows) == 2
            assert rows[1][rows[0].index("学号")] == empty

            # New writes must reject blank course codes; do not weaken that contract
            # just to construct historical data. Selection writes also normalize ""
            # to "未选", so genuine legacy empty strings require isolated PG seeding.
            rejected = request("POST", "/data/courses", expected=422, json={
                "year": AUDIT_ACADEMIC_YEAR, "term": scope["term"], "grade": "高一",
                "course_code": "", "course_name": "历史空代码夹具", "quota": 3,
            }).json()
            assert any(error["loc"] == ["body", "course_code"] for error in rejected["detail"])
            legacy, outside = "LEGACY_" + tag, "OUTSIDE_" + tag
            for no, grade in [(legacy, "高一"), (outside, "高二")]:
                request("POST", "/data/students", json={**scope, "grade": grade, "student_no": no, "name": no})
            seeded = client.post(base + "/__audit__/xbk/legacy-empty-codes", json={
                "year": AUDIT_ACADEMIC_YEAR, "term": scope["term"], "tag": tag,
            })
            assert seeded.status_code == 200, (seeded.status_code, seeded.text[:800])
            fixture = seeded.json()
            blank_course_id = fixture["course"]["id"]
            assert fixture["status"] == "isolated-xbk-fixture-seeded"
            assert fixture["year"] == AUDIT_ACADEMIC_YEAR and fixture["term"] == scope["term"]
            assert fixture["class_name"] == scope["class_name"]
            assert isinstance(blank_course_id, int) and blank_course_id > 0
            assert fixture["course"]["course_code"] == ""
            assert len(fixture["selections"]) == 2
            assert {(row["student_no"], row["grade"], row["course_code"])
                    for row in fixture["selections"]} == {
                (legacy, "高二", ""), (outside, "高一", ""),
            }
            stats = request("GET", "/analysis/course-stats", params=scope).json()["items"]
            unselected = [row for row in stats if row["course_code"] in {"", "未选"}]
            assert len(unselected) == 1 and unselected[0]["course_code"] == "未选"
            assert unselected[0]["count"] == 2
            exported = request("GET", "/export", params={**scope, "scope": "unselected"})
            rows = list(load_workbook(BytesIO(exported.content)).active.values)
            assert {row[rows[0].index("学号")] for row in rows[1:]} == {empty, legacy}

            # Same student number in another semester must remain suspended there.
            for no, grade in [(chosen, "高一"), (outside, "高二")]:
                request("POST", "/data/students", json={
                    **scope, "term": "下学期", "grade": grade, "student_no": no, "name": no,
                })
            for semester in [None, "上学期", "下学期"]:
                filters = {key: value for key, value in scope.items() if key != "term"}
                if semester:
                    filters["term"] = semester
                exported = request("GET", "/export", params={**filters, "scope": "suspended"})
                rows = list(load_workbook(BytesIO(exported.content)).active.values)
                records = [dict(zip(rows[0], row)) for row in rows[1:]]
                expected = [] if semester == "上学期" else [(AUDIT_ACADEMIC_YEAR, "下学期", chosen)]
                assert [(r["学年"], r["学期"], r["学号"]) for r in records] == expected
        finally:
            # Cleanup is restricted to this unique 2032-2033 class, even on assertion failure.
            for semester in ["上学期", "下学期"]:
                request("DELETE", "/data", params={
                    "year": AUDIT_ACADEMIC_YEAR, "term": semester, "class_name": scope["class_name"], "scope": "students",
                })
            for fixture_course_id in [course_id, blank_course_id]:
                if fixture_course_id is not None:
                    request("DELETE", f"/data/courses/{fixture_course_id}", expected=204)
            remaining = request("GET", "/data/students", params=scope).json()
            assert remaining["total"] == 0
    print("\n2032-2033 multipart blank-selection import: summary/list/stats/export consistent")
