import io
from typing import List

import pandas as pd
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def style_worksheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="F0F2F5")
    header_font = Font(bold=True, color="000000")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_align = Alignment(vertical="center", wrap_text=True)
    max_row, max_col = ws.max_row, ws.max_column
    if max_row < 1 or max_col < 1:
        return
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"
    ws.row_dimensions[1].height = 22
    for col in range(1, max_col + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill, cell.font, cell.alignment = header_fill, header_font, header_align
    for row in range(2, max_row + 1):
        ws.row_dimensions[row].height = 18
        for col in range(1, max_col + 1):
            ws.cell(row=row, column=col).alignment = body_align
    for col in range(1, max_col + 1):
        max_len = max(
            (len(str(ws.cell(row=row, column=col).value))
             for row in range(1, min(max_row, 200) + 1)
             if ws.cell(row=row, column=col).value is not None),
            default=0,
        )
        ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 10), 48)


def template_columns(scope: str) -> List[str]:
    if scope == "students":
        return ["学年", "学期", "年级", "班级", "学号", "姓名", "性别"]
    if scope == "courses":
        return ["学年", "学期", "年级", "课程代码", "课程名称", "课程负责人", "各班限报人数", "上课地点"]
    return ["学年", "学期", "年级", "学号", "姓名", "课程代码"]


def build_template_response(scope: str) -> StreamingResponse:
    columns = template_columns(scope)
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(columns=columns).to_excel(writer, index=False, sheet_name="template")
        notes = [
            ["项目", "说明"],
            ["用途", f"XBK {scope} 导入模板；请仅在 template 工作表填入正式数据。"],
            ["学年", "必须使用 xxxx-xxxx 格式，例如 2026-2027；不要只填写 2026。"],
            ["学期", "例如：上学期。"],
            ["学号", "按文本填写并保证全校唯一；高一、高二不要直接使用相同学号，建议使用 G1-101、G2-101。"],
            ["课程代码", "按文本填写并保证同一学年、学期内唯一；建议使用 G1-09、G2-09。学生选课结果中的课程代码必须存在于课程目录。"],
            ["各班限报人数", "课程目录字段必须是非负整数；原始字段“课程人数”也可被识别为该字段。"],
            ["导入顺序", "学生名单 → 课程目录 → 学生选课结果。"],
            ["示例", "高一学号 G1-101；高二学号 G2-101；高一课程 G1-09；高二课程 G2-09。"],
            ["注意", "示例与说明工作表不会被导入；请勿修改 template 工作表表头。"],
        ]
        pd.DataFrame(notes[1:], columns=notes[0]).to_excel(writer, index=False, sheet_name="示例与说明")
        template_ws, notes_ws = writer.book["template"], writer.book["示例与说明"]
        style_worksheet(template_ws)
        style_worksheet(notes_ws)
        for ws in (template_ws, notes_ws):
            for cell in ws[1]:
                cell.number_format = "@"
        for index, header in enumerate(columns, start=1):
            if header in {"学年", "学期", "年级", "学号", "课程代码"}:
                for row in range(2, 1002):
                    template_ws.cell(row=row, column=index).number_format = "@"
        notes_ws.column_dimensions["A"].width = 20
        notes_ws.column_dimensions["B"].width = 100
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="xbk_{scope}_template.xlsx"'},
    )
