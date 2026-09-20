"""小组讨论管理端路由（/admin/*）。

路由路径写全（不含 prefix），由 `group_discussion.py` 统一聚合。
模块间仍保持 public-config、用户端、管理端的顺序；管理端内部路由的排列不作为 API
合同，实际路径、方法、响应模型和权限依赖才是兼容性边界。
"""
import io
from datetime import date, datetime
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession
from starlette.concurrency import run_in_threadpool

from app.core.deps import get_db, require_admin
from app.services import classroom as svc
from app.schemas.agents import (
    GroupDiscussionAdminSessionListResponse,
    GroupDiscussionAdminSessionOut,
    GroupDiscussionAdminAnalyzeRequest,
    GroupDiscussionAdminAnalyzeResponse,
    GroupDiscussionAdminCompareAnalyzeRequest,
    GroupDiscussionMemberOut,
    GroupDiscussionAdminMemberListResponse,
    GroupDiscussionAdminDeleteSessionsRequest,
    GroupDiscussionStudentProfileRequest,
    GroupDiscussionCrossSystemRequest,
)
from app.services.agents.group_discussion import (
    admin_analyze_session,
    admin_compare_analyze_sessions,
    admin_cross_system_analysis,
    admin_student_profile_analysis,
    admin_delete_session,
    admin_delete_sessions,
    admin_list_sessions,
    admin_list_members,
    admin_list_all_sessions,
    list_classes,
)

router = APIRouter()


@router.get("/admin/sessions", response_model=GroupDiscussionAdminSessionListResponse)
async def admin_get_sessions(
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    class_name: Optional[str] = Query(None),
    group_no: Optional[str] = Query(None),
    group_name: Optional[str] = Query(None),
    user_name: Optional[str] = Query(None),
    keyword: Optional[str] = Query(None, description="关键词搜索（班级/组号/组名/姓名）"),
    page: int = Query(1, ge=1),
    size: int = Query(20, ge=1, le=200),
    db: AsyncSession = Depends(get_db),
    _: Dict[str, Any] = Depends(require_admin),
) -> GroupDiscussionAdminSessionListResponse:
    rows, total, page_n, total_pages = await admin_list_sessions(
        db,
        start_date=start_date,
        end_date=end_date,
        class_name=class_name,
        group_no=group_no,
        group_name=group_name,
        user_name=user_name,
        keyword=keyword,
        page=page,
        size=size,
    )
    items = [
        GroupDiscussionAdminSessionOut(
            id=int(r.id),
            session_date=r.session_date,
            class_name=str(r.class_name),
            group_no=str(r.group_no),
            group_name=(str(r.group_name).strip() if r.group_name else None),
            message_count=int(r.message_count or 0),
            created_at=r.created_at,
            last_message_at=r.last_message_at,
        )
        for r in rows
    ]
    return GroupDiscussionAdminSessionListResponse(
        items=items,
        total=total,
        page=page_n,
        page_size=size,
        total_pages=total_pages,
    )


@router.get("/admin/export-sessions")
async def admin_export_sessions(
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    class_name: Optional[str] = Query(None),
    group_no: Optional[str] = Query(None),
    group_name: Optional[str] = Query(None),
    user_name: Optional[str] = Query(None),
    keyword: Optional[str] = Query(None, description="关键词搜索（班级/组号/组名/姓名）"),
    limit: int = Query(5000, ge=1, le=10000, description="最大导出行数，默认 5000，上限 10000"),
    db: AsyncSession = Depends(get_db),
    _: Dict[str, Any] = Depends(require_admin),
):
    """导出筛选后的会话列表为 Excel"""
    export_limit = min(int(limit or 5000), 10000)
    rows = await admin_list_all_sessions(
        db,
        start_date=start_date,
        end_date=end_date,
        class_name=class_name,
        group_no=group_no,
        group_name=group_name,
        user_name=user_name,
        keyword=keyword,
        limit=export_limit,
    )
    if len(rows) > export_limit:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"导出结果超过 {export_limit} 条，请缩小筛选范围后重试",
        )

    def _build_xlsx(sessions) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "小组讨论会话"

        headers = ["ID", "日期", "班级", "组号", "组名", "消息数", "创建时间", "最后消息"]
        header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        header_font = Font(bold=True)
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align

        for row_idx, s in enumerate(sessions, 2):
            ws.cell(row=row_idx, column=1, value=s.id)
            ws.cell(row=row_idx, column=2, value=str(s.session_date) if s.session_date else "")
            ws.cell(row=row_idx, column=3, value=s.class_name)
            ws.cell(row=row_idx, column=4, value=s.group_no)
            ws.cell(row=row_idx, column=5, value=s.group_name or "")
            ws.cell(row=row_idx, column=6, value=int(s.message_count or 0))
            ws.cell(row=row_idx, column=7, value=s.created_at.strftime("%Y-%m-%d %H:%M") if s.created_at else "")
            ws.cell(row=row_idx, column=8, value=s.last_message_at.strftime("%Y-%m-%d %H:%M") if s.last_message_at else "")

        col_widths = [8, 14, 18, 10, 20, 10, 18, 18]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

        ws.freeze_panes = "A2"

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    excel_bytes = await run_in_threadpool(_build_xlsx, rows)
    now_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    return Response(
        content=excel_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="group_discussion_sessions_{now_str}.xlsx"'
        },
    )


@router.delete("/admin/sessions/{session_id}", status_code=status.HTTP_200_OK)
async def admin_delete_session_api(
    session_id: int,
    db: AsyncSession = Depends(get_db),
    _: Dict[str, Any] = Depends(require_admin),
) -> Any:
    await admin_delete_session(db, session_id=session_id)

    # 发布事件
    await svc.publish("admin_global", {"type": "discussion_changed", "action": "delete", "id": session_id})

    return {"success": True}


@router.post("/admin/sessions/batch-delete", status_code=status.HTTP_200_OK)
async def admin_batch_delete_sessions_api(
    payload: GroupDiscussionAdminDeleteSessionsRequest,
    db: AsyncSession = Depends(get_db),
    _: Dict[str, Any] = Depends(require_admin),
) -> Any:
    deleted = await admin_delete_sessions(db, session_ids=payload.session_ids)
    await svc.publish("admin_global", {"type": "discussion_changed", "action": "batch_delete"})
    return {"success": True, "deleted": deleted}


@router.get("/admin/members", response_model=GroupDiscussionAdminMemberListResponse)
async def admin_get_members(
    session_id: int = Query(..., ge=1),
    db: AsyncSession = Depends(get_db),
    _: Dict[str, Any] = Depends(require_admin),
) -> GroupDiscussionAdminMemberListResponse:
    members = await admin_list_members(db, session_id=session_id)
    items = []
    for m in members:
        items.append(
            GroupDiscussionMemberOut(
                user_id=m.user_id,
                username=getattr(m.user, "username", None),
                full_name=getattr(m.user, "full_name", None),
                student_id=getattr(m.user, "student_id", None),
                joined_at=m.joined_at,
                muted_until=m.muted_until,
            )
        )
    return GroupDiscussionAdminMemberListResponse(items=items)


@router.get("/admin/classes", response_model=List[str])
async def admin_list_classes(
    date: Optional[date] = Query(None),
    db: AsyncSession = Depends(get_db),
    _: Dict[str, Any] = Depends(require_admin),
) -> List[str]:
    return await list_classes(db, date=date)


@router.post("/admin/analyze", response_model=GroupDiscussionAdminAnalyzeResponse)
async def admin_analyze(
    payload: GroupDiscussionAdminAnalyzeRequest,
    db: AsyncSession = Depends(get_db),
    admin_user: Dict[str, Any] = Depends(require_admin),
) -> GroupDiscussionAdminAnalyzeResponse:
    r = await admin_analyze_session(
        db,
        session_id=payload.session_id,
        agent_id=payload.agent_id,
        admin_user=admin_user,
        analysis_type=payload.analysis_type,
        prompt=payload.prompt,
    )
    return GroupDiscussionAdminAnalyzeResponse(
        analysis_id=int(r.id),
        result_text=str(r.result_text),
        created_at=r.created_at,
    )


@router.post("/admin/compare-analyze", response_model=GroupDiscussionAdminAnalyzeResponse)
async def admin_compare_analyze(
    payload: GroupDiscussionAdminCompareAnalyzeRequest,
    db: AsyncSession = Depends(get_db),
    admin_user: Dict[str, Any] = Depends(require_admin),
) -> GroupDiscussionAdminAnalyzeResponse:
    r = await admin_compare_analyze_sessions(
        db,
        session_ids=payload.session_ids,
        agent_id=payload.agent_id,
        admin_user=admin_user,
        bucket_seconds=payload.bucket_seconds,
        analysis_type=payload.analysis_type,
        prompt=payload.prompt,
        use_cache=payload.use_cache,
    )
    return GroupDiscussionAdminAnalyzeResponse(
        analysis_id=int(r.id),
        result_text=str(r.result_text),
        created_at=r.created_at,
    )


@router.post("/admin/student-profile", response_model=GroupDiscussionAdminAnalyzeResponse)
async def admin_student_profile(
    payload: GroupDiscussionStudentProfileRequest,
    db: AsyncSession = Depends(get_db),
    admin_user: Dict[str, Any] = Depends(require_admin),
) -> GroupDiscussionAdminAnalyzeResponse:
    r = await admin_student_profile_analysis(
        db,
        session_id=payload.session_id,
        user_id=payload.user_id,
        agent_id=payload.agent_id,
        admin_user=admin_user,
    )
    return GroupDiscussionAdminAnalyzeResponse(
        analysis_id=int(r.id),
        result_text=str(r.result_text),
        created_at=r.created_at,
    )


@router.post("/admin/cross-system-analyze", response_model=GroupDiscussionAdminAnalyzeResponse)
async def admin_cross_system_analyze(
    payload: GroupDiscussionCrossSystemRequest,
    db: AsyncSession = Depends(get_db),
    admin_user: Dict[str, Any] = Depends(require_admin),
) -> GroupDiscussionAdminAnalyzeResponse:
    r = await admin_cross_system_analysis(
        db,
        session_ids=payload.session_ids,
        agent_id=payload.agent_id,
        admin_user=admin_user,
        target_date=payload.date,
        class_name=payload.class_name,
    )
    return GroupDiscussionAdminAnalyzeResponse(
        analysis_id=int(r.id),
        result_text=str(r.result_text),
        created_at=r.created_at,
    )
