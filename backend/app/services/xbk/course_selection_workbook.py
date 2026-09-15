"""Versioned, DB-free parser for XBK all-class course-selection workbooks.

The parser deliberately stops before any database write.  It validates the
exported workbook as a complete roster snapshot and returns an explicit input
contract for the transactional selection service that owns final authorization,
current-roster checks and capacity arbitration.
"""
from __future__ import annotations
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from hashlib import sha256
from io import BytesIO
import json
from pathlib import PurePosixPath
import stat
from typing import Iterable, Mapping, Optional, Sequence
from zipfile import BadZipFile, ZIP_DEFLATED, ZIP_STORED, ZipFile
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

WORKBOOK_FORMAT = 'xbk-course-selection'
WORKBOOK_VERSION = '1'
METADATA_SHEET = '__XBK工作簿信息'
CATALOG_SHEET = '校本课程目录'
INSTRUCTION = '只填写本班课程代码；剩余名额仅反映本文件，提交以服务器校验为准。'
UNSELECTED_CODE = '未选'
MAX_XLSX_ARCHIVE_ENTRIES = 1024
MAX_XLSX_WORKSHEETS = 256
MAX_XLSX_TOTAL_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
MAX_XLSX_MEMBER_UNCOMPRESSED_BYTES = 32 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 250
_ALLOWED_XLSX_COMPRESSION = frozenset((ZIP_STORED, ZIP_DEFLATED))


class WorkbookArchiveError(ValueError):
    """Structured rejection raised before untrusted XLSX reaches openpyxl."""

    def __init__(self, code: str, message: str, *, status_code: int) -> None:
        self.code = code
        self.message = message
        self.status_code = status_code
        super().__init__(message)

    def to_dict(self) -> dict:
        return {'code': self.code, 'message': self.message, 'issues': []}


def _archive_invalid(message: str) -> WorkbookArchiveError:
    return WorkbookArchiveError('XBK_WORKBOOK_ARCHIVE_INVALID', message, status_code=422)


def _archive_too_large(message: str) -> WorkbookArchiveError:
    return WorkbookArchiveError('XBK_WORKBOOK_ARCHIVE_TOO_LARGE', message, status_code=413)


def _read_xlsx_archive_members(content: bytes) -> list:
    try:
        with ZipFile(BytesIO(content)) as archive:
            return archive.infolist()
    except (BadZipFile, OSError, ValueError) as exc:
        raise _archive_invalid('文件不是有效的 XLSX ZIP 容器') from exc


def _validate_archive_entry_count(members: Sequence[object]) -> None:
    if len(members) > MAX_XLSX_ARCHIVE_ENTRIES:
        raise _archive_too_large(f'工作簿 ZIP 条目不能超过 {MAX_XLSX_ARCHIVE_ENTRIES} 个')


def _archive_member_path_is_unsafe(name: str, path: PurePosixPath) -> bool:
    if not name or '\x00' in name or '\\' in name:
        return True
    if path.is_absolute() or '..' in path.parts:
        return True
    return bool(path.parts and ':' in path.parts[0])


def _validate_archive_member_identity(member, seen_names: set[str]) -> str:
    path = PurePosixPath(member.filename)
    unix_mode = member.external_attr >> 16 & 61440
    if _archive_member_path_is_unsafe(member.filename, path) or unix_mode == stat.S_IFLNK:
        raise _archive_invalid('工作簿 ZIP 包含不安全的成员路径或符号链接')
    normalized_name = path.as_posix()
    if normalized_name in seen_names:
        raise _archive_invalid('工作簿 ZIP 包含重复成员')
    seen_names.add(normalized_name)
    return normalized_name


def _validate_archive_member_format(member) -> None:
    if member.flag_bits & 1:
        raise _archive_invalid('工作簿 ZIP 不能包含加密成员')
    if not member.is_dir() and member.compress_type not in _ALLOWED_XLSX_COMPRESSION:
        raise _archive_invalid('工作簿 ZIP 使用了不支持的压缩方式')
    if member.file_size < 0 or member.compress_size < 0:
        raise _archive_invalid('工作簿 ZIP 成员大小信息无效')


def _validate_archive_member_budget(member) -> None:
    if member.file_size > MAX_XLSX_MEMBER_UNCOMPRESSED_BYTES:
        raise _archive_too_large(f'工作簿 ZIP 单个条目解压后不能超过 {MAX_XLSX_MEMBER_UNCOMPRESSED_BYTES // (1024 * 1024)} MiB')
    if member.file_size and member.compress_size == 0:
        raise _archive_too_large('工作簿 ZIP 成员压缩比异常')
    if member.compress_size and member.file_size / member.compress_size > MAX_XLSX_COMPRESSION_RATIO:
        raise _archive_too_large(f'工作簿 ZIP 单个条目压缩比不能超过 {MAX_XLSX_COMPRESSION_RATIO}:1')


def _is_worksheet_member(normalized_name: str, member) -> bool:
    return normalized_name.startswith('xl/worksheets/') and normalized_name.endswith('.xml') and (not member.is_dir())


def _archive_totals(members: Sequence[object]) -> tuple[int, int, int]:
    total_compressed = 0
    total_uncompressed = 0
    worksheet_count = 0
    seen_names: set[str] = set()
    for member in members:
        normalized_name = _validate_archive_member_identity(member, seen_names)
        _validate_archive_member_format(member)
        _validate_archive_member_budget(member)
        total_compressed += member.compress_size
        total_uncompressed += member.file_size
        if total_uncompressed > MAX_XLSX_TOTAL_UNCOMPRESSED_BYTES:
            raise _archive_too_large(f'工作簿 ZIP 解压后总大小不能超过 {MAX_XLSX_TOTAL_UNCOMPRESSED_BYTES // (1024 * 1024)} MiB')
        worksheet_count += int(_is_worksheet_member(normalized_name, member))
    return (total_compressed, total_uncompressed, worksheet_count)


def _validate_archive_totals(total_compressed: int, total_uncompressed: int, worksheet_count: int) -> None:
    excessive_ratio = total_compressed == 0 or total_uncompressed / total_compressed > MAX_XLSX_COMPRESSION_RATIO
    if total_uncompressed and excessive_ratio:
        raise _archive_too_large(f'工作簿 ZIP 总压缩比不能超过 {MAX_XLSX_COMPRESSION_RATIO}:1')
    if worksheet_count > MAX_XLSX_WORKSHEETS:
        raise _archive_too_large(f'工作簿工作表不能超过 {MAX_XLSX_WORKSHEETS} 个')


def validate_xlsx_archive_budget(content: bytes) -> None:
    """Statically bound an XLSX ZIP before openpyxl decompresses any member."""
    members = _read_xlsx_archive_members(content)
    _validate_archive_entry_count(members)
    _validate_archive_totals(*_archive_totals(members))


@dataclass(frozen=True)


class WorkbookCourse:
    course_code: str
    course_name: str
    quota: int


@dataclass(frozen=True)


class WorkbookStudent:
    sheet: str
    grade: str
    class_name: str
    student_no: str
    name: str
    baseline_course: Optional[str]


@dataclass(frozen=True)


class WorkbookExpectation:
    """Server-held export baseline used to reject stale or substituted files."""
    year: str
    term: str
    baseline_id: str
    students: Sequence[WorkbookStudent]
    courses: Sequence[WorkbookCourse]


@dataclass(frozen=True)


class WorkbookSelectionInput:
    """One complete-roster row for the future R1 transactional execute service."""
    sheet: str
    row: int
    year: str
    term: str
    grade: str
    class_name: str
    student_no: str
    name: str
    baseline_course: Optional[str]
    submitted_course: Optional[str]
    action: str


@dataclass(frozen=True)


class WorkbookSuggestion:
    course_code: str
    course_name: str
    remaining: int


@dataclass(frozen=True)


class WorkbookIssue:
    sheet: str
    row: Optional[int]
    student_no: Optional[str]
    course_code: Optional[str]
    reason: str
    suggestions: Sequence[WorkbookSuggestion] = field(default_factory=tuple)

    def to_dict(self) -> dict:
        return asdict(self)


class WorkbookValidationError(ValueError):

    def __init__(self, issues: Sequence[WorkbookIssue]):
        self.issues = tuple(issues)
        super().__init__(f'选课工作簿包含 {len(self.issues)} 个冲突')

    def to_dict(self) -> dict:
        return {'code': 'XBK_WORKBOOK_INVALID', 'issues': [issue.to_dict() for issue in self.issues]}


@dataclass(frozen=True)


class ParsedCourseSelectionWorkbook:
    format: str
    version: str
    year: str
    term: str
    baseline_id: str
    rows: Sequence[WorkbookSelectionInput]


def normalize_cell(value: object) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _manifest_payload(year: str, term: str, students: Iterable[WorkbookStudent], courses: Iterable[WorkbookCourse]) -> dict:
    return {'format': WORKBOOK_FORMAT, 'version': WORKBOOK_VERSION, 'year': year, 'term': term, 'students': [{'sheet': item.sheet, 'grade': item.grade, 'class_name': item.class_name, 'student_no': item.student_no, 'name': item.name, 'baseline_course': item.baseline_course} for item in students], 'courses': [{'course_code': item.course_code, 'course_name': item.course_name, 'quota': item.quota} for item in courses]}


def compute_baseline_id(year: str, term: str, students: Iterable[WorkbookStudent], courses: Iterable[WorkbookCourse]) -> str:
    payload = json.dumps(_manifest_payload(year, term, students, courses), ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode('utf-8')
    return sha256(payload).hexdigest()


def write_metadata_sheet(ws: Worksheet, *, year: str, term: str, students: Sequence[WorkbookStudent], courses: Sequence[WorkbookCourse]) -> str:
    """Write a row-oriented manifest so large rosters never exceed one cell."""
    baseline_id = compute_baseline_id(year, term, students, courses)
    ws.append(['格式', WORKBOOK_FORMAT])
    ws.append(['版本', WORKBOOK_VERSION])
    ws.append(['学年', year])
    ws.append(['学期', term])
    ws.append(['导出基线', baseline_id])
    ws.append([])
    ws.append(['类型', '工作表', '年级', '班级', '学号/课程代码', '姓名/课程名称', '基线课程', '限额'])
    for student in students:
        ws.append(['学生', student.sheet, student.grade, student.class_name, student.student_no, student.name, student.baseline_course, None])
    for course in courses:
        ws.append(['课程', None, None, None, course.course_code, course.course_name, None, course.quota])
    return baseline_id


def _metadata_issue(reason: str, row: Optional[int]=None) -> WorkbookIssue:
    return WorkbookIssue(METADATA_SHEET, row, None, None, reason)


def _read_metadata_identity(ws: Worksheet, issues: list[WorkbookIssue]) -> tuple[str, str, str]:
    keys = {normalize_cell(ws.cell(row=row, column=1).value): normalize_cell(ws.cell(row=row, column=2).value) for row in range(1, 6)}
    if keys.get('格式') != WORKBOOK_FORMAT:
        issues.append(_metadata_issue('工作簿格式标识无效', 1))
    if keys.get('版本') != WORKBOOK_VERSION:
        issues.append(_metadata_issue('工作簿版本不受支持', 2))
    year = keys.get('学年') or ''
    term = keys.get('学期') or ''
    baseline_id = keys.get('导出基线') or ''
    for value, reason, row in ((year, '缺少导出学年', 3), (term, '缺少导出学期', 4), (baseline_id, '缺少导出基线', 5)):
        if not value:
            issues.append(_metadata_issue(reason, row))
    return (year, term, baseline_id)


def _validate_metadata_header(ws: Worksheet, issues: list[WorkbookIssue]) -> None:
    expected_header = ['类型', '工作表', '年级', '班级', '学号/课程代码', '姓名/课程名称', '基线课程', '限额']
    actual_header = [normalize_cell(ws.cell(row=7, column=col).value) for col in range(1, 9)]
    if actual_header != expected_header:
        issues.append(_metadata_issue('工作簿元数据表头已被修改', 7))


def _read_metadata_student(ws: Worksheet, row: int, issues: list[WorkbookIssue]) -> Optional[WorkbookStudent]:
    values = [normalize_cell(ws.cell(row=row, column=col).value) for col in range(2, 8)]
    sheet, grade, class_name, student_no, name, baseline_course = values
    if not all((sheet, grade, class_name, student_no, name)):
        issues.append(_metadata_issue('学生基线信息不完整', row))
        return None
    return WorkbookStudent(sheet, grade, class_name, student_no, name, baseline_course)


def _parse_metadata_quota(value: object) -> Optional[int]:
    try:
        quota = int(value)
    except (TypeError, ValueError):
        return None
    if quota < 0 or quota != value:
        return None
    return quota


def _read_metadata_course(ws: Worksheet, row: int, issues: list[WorkbookIssue]) -> Optional[WorkbookCourse]:
    code = normalize_cell(ws.cell(row=row, column=5).value)
    name = normalize_cell(ws.cell(row=row, column=6).value)
    quota = _parse_metadata_quota(ws.cell(row=row, column=8).value)
    if quota is None:
        issues.append(_metadata_issue('课程限额无效', row))
        return None
    if not code or not name:
        issues.append(_metadata_issue('课程基线信息不完整', row))
        return None
    return WorkbookCourse(code, name, quota)


def _read_metadata_records(ws: Worksheet, issues: list[WorkbookIssue]) -> tuple[list[WorkbookStudent], list[WorkbookCourse]]:
    students: list[WorkbookStudent] = []
    courses: list[WorkbookCourse] = []
    for row in range(8, ws.max_row + 1):
        kind = normalize_cell(ws.cell(row=row, column=1).value)
        if not kind:
            continue
        if kind == '学生':
            student = _read_metadata_student(ws, row, issues)
            if student is not None:
                students.append(student)
            continue
        if kind == '课程':
            course = _read_metadata_course(ws, row, issues)
            if course is not None:
                courses.append(course)
            continue
        issues.append(_metadata_issue('工作簿元数据包含未知记录类型', row))
    return (students, courses)


def _read_metadata(ws: Worksheet) -> tuple[str, str, str, list[WorkbookStudent], list[WorkbookCourse], list[WorkbookIssue]]:
    issues: list[WorkbookIssue] = []
    year, term, baseline_id = _read_metadata_identity(ws, issues)
    _validate_metadata_header(ws, issues)
    students, courses = _read_metadata_records(ws, issues)
    return (year, term, baseline_id, students, courses, issues)


def _suggestions_for_class(class_key: tuple[str, str], courses: Sequence[WorkbookCourse], occupancy: Mapping[tuple[str, str, str], int], *, limit: int=3) -> tuple[WorkbookSuggestion, ...]:
    grade, class_name = class_key
    result = []
    for course in courses:
        remaining = course.quota - occupancy.get((grade, class_name, course.course_code), 0)
        if remaining > 0:
            result.append(WorkbookSuggestion(course.course_code, course.course_name, remaining))
    result.sort(key=lambda item: (-item.remaining, item.course_code))
    return tuple(result[:limit])


def _validate_catalog(ws: Optional[Worksheet], courses: Sequence[WorkbookCourse]) -> list[WorkbookIssue]:
    if ws is None:
        return [WorkbookIssue(CATALOG_SHEET, None, None, None, '缺少课程目录工作表')]
    issues: list[WorkbookIssue] = []
    if not ws.protection.sheet:
        issues.append(WorkbookIssue(CATALOG_SHEET, None, None, None, '课程目录保护已被移除'))
    expected_header = ['课程代码', '课程名称', '课程负责人', '各班限报人数', '上课地点']
    actual_header = [normalize_cell(ws.cell(row=2, column=col).value) for col in range(1, 6)]
    if actual_header != expected_header:
        issues.append(WorkbookIssue(CATALOG_SHEET, 2, None, None, '课程目录表头已被修改'))
    expected = {course.course_code: course for course in courses}
    seen: set[str] = set()
    for row in range(3, ws.max_row + 1):
        code = normalize_cell(ws.cell(row=row, column=1).value)
        if not code:
            continue
        name = normalize_cell(ws.cell(row=row, column=2).value)
        quota_value = ws.cell(row=row, column=4).value
        course = expected.get(code)
        if course is None:
            issues.append(WorkbookIssue(CATALOG_SHEET, row, None, code, '课程目录包含未知课程'))
            continue
        if code in seen:
            issues.append(WorkbookIssue(CATALOG_SHEET, row, None, code, '课程目录包含重复课程'))
            continue
        seen.add(code)
        try:
            quota = int(quota_value)
        except (TypeError, ValueError):
            quota = -1
        if name != course.course_name or quota != course.quota:
            issues.append(WorkbookIssue(CATALOG_SHEET, row, None, code, '课程目录与导出基线不一致'))
    for code in sorted(set(expected) - seen):
        issues.append(WorkbookIssue(CATALOG_SHEET, None, None, code, '课程目录缺少导出课程'))
    return issues


def read_course_selection_workbook_expectation(content: bytes | BytesIO) -> WorkbookExpectation:
    """Recover the immutable export contract embedded in one workbook.

    This is intentionally self-contained: it only trusts metadata whose digest
    recomputes correctly. Callers must still parse the full workbook against
    the returned expectation and bind the resulting plan to a server-signed
    preview token before applying any database change.
    """
    raw_content = content.getvalue() if hasattr(content, 'getvalue') else bytes(content)
    validate_xlsx_archive_budget(raw_content)
    try:
        wb = load_workbook(BytesIO(raw_content), data_only=False, read_only=True)
    except Exception as exc:
        raise WorkbookValidationError([WorkbookIssue('工作簿', None, None, None, '文件不是有效的 XLSX 工作簿')]) from exc
    try:
        if METADATA_SHEET not in wb.sheetnames:
            raise WorkbookValidationError([_metadata_issue('缺少版本化工作簿元数据')])
        metadata_ws = wb[METADATA_SHEET]
        year, term, baseline_id, students, courses, issues = _read_metadata(metadata_ws)
        if metadata_ws.sheet_state != 'veryHidden':
            issues.append(_metadata_issue('工作簿元数据保护状态已被修改'))
        computed = compute_baseline_id(year, term, students, courses) if year and term else ''
        if baseline_id and baseline_id != computed:
            issues.append(_metadata_issue('导出基线校验失败，工作簿结构或名册已被修改', 5))
        if issues:
            raise WorkbookValidationError(issues)
        return WorkbookExpectation(year=year, term=term, baseline_id=baseline_id, students=tuple(students), courses=tuple(courses))
    finally:
        wb.close()


def _workbook_bytes(content: bytes | BytesIO) -> bytes:
    return content.getvalue() if hasattr(content, 'getvalue') else bytes(content)


def _load_course_selection_workbook(raw_content: bytes, *, read_only: bool=False):
    try:
        return load_workbook(BytesIO(raw_content), data_only=False, read_only=read_only)
    except Exception as exc:
        raise WorkbookValidationError([WorkbookIssue('工作簿', None, None, None, '文件不是有效的 XLSX 工作簿')]) from exc


def _read_validated_metadata(wb):
    if METADATA_SHEET not in wb.sheetnames:
        raise WorkbookValidationError([_metadata_issue('缺少版本化工作簿元数据')])
    metadata_ws = wb[METADATA_SHEET]
    year, term, baseline_id, students, courses, issues = _read_metadata(metadata_ws)
    if metadata_ws.sheet_state != 'veryHidden':
        issues.append(_metadata_issue('工作簿元数据保护状态已被修改'))
    computed = compute_baseline_id(year, term, students, courses) if year and term else ''
    if baseline_id and baseline_id != computed:
        issues.append(_metadata_issue('导出基线校验失败，工作簿结构或名册已被修改', 5))
    return (year, term, baseline_id, students, courses, issues)


def _validate_expected_contract(year: str, term: str, baseline_id: str, students: Sequence[WorkbookStudent], courses: Sequence[WorkbookCourse], expected: Optional[WorkbookExpectation]) -> list[WorkbookIssue]:
    if expected is None:
        return []
    issues: list[WorkbookIssue] = []
    if (year, term) != (expected.year, expected.term):
        issues.append(_metadata_issue('工作簿学年或学期与当前导入范围不一致'))
    if baseline_id != expected.baseline_id:
        issues.append(_metadata_issue('工作簿导出基线已过期，请重新导出'))
    expected_digest = compute_baseline_id(expected.year, expected.term, expected.students, expected.courses)
    if expected_digest != expected.baseline_id:
        issues.append(_metadata_issue('服务器导出基线合同无效'))
    if list(students) != list(expected.students) or list(courses) != list(expected.courses):
        issues.append(_metadata_issue('工作簿名册或课程基线与服务器记录不一致'))
    return issues


def _validate_workbook_sheet_set(wb, expected_sheets: set[str]) -> list[WorkbookIssue]:
    issues: list[WorkbookIssue] = []
    allowed_sheets = expected_sheets | {CATALOG_SHEET, METADATA_SHEET}
    for name in wb.sheetnames:
        if name not in allowed_sheets:
            issues.append(WorkbookIssue(name, None, None, None, '包含未知或复制的工作表'))
    for name in sorted(expected_sheets - set(wb.sheetnames)):
        issues.append(WorkbookIssue(name, None, None, None, '缺少导出的班级工作表'))
    return issues


def _build_roster_by_sheet(students: Sequence[WorkbookStudent], issues: list[WorkbookIssue]) -> dict[str, dict[str, WorkbookStudent]]:
    roster_by_sheet: dict[str, dict[str, WorkbookStudent]] = defaultdict(dict)
    for student in students:
        if student.student_no in roster_by_sheet[student.sheet]:
            issues.append(WorkbookIssue(student.sheet, None, student.student_no, None, '导出基线内学生重复'))
        roster_by_sheet[student.sheet][student.student_no] = student
    return roster_by_sheet


def _validate_class_sheet_structure(ws: Worksheet, sheet: str) -> list[WorkbookIssue]:
    issues: list[WorkbookIssue] = []
    if ws.sheet_state != 'visible':
        issues.append(WorkbookIssue(sheet, None, None, None, '班级工作表被隐藏'))
    if not ws.protection.sheet:
        issues.append(WorkbookIssue(sheet, None, None, None, '班级工作表保护已被移除'))
    header = [normalize_cell(ws.cell(row=1, column=col).value) for col in range(1, 5)]
    if header != ['班级', '学号', '姓名', '课程代码']:
        issues.append(WorkbookIssue(sheet, 1, None, None, '班级工作表表头已被修改'))
    if normalize_cell(ws.cell(row=1, column=6).value) != INSTRUCTION:
        issues.append(WorkbookIssue(sheet, 1, None, None, '填写说明已被修改'))
    summary_header = [normalize_cell(ws.cell(row=2, column=col).value) for col in range(6, 11)]
    if summary_header != ['课程代码', '课程名称', '本班限额', '已选', '剩余']:
        issues.append(WorkbookIssue(sheet, 2, None, None, '本班剩余名额表头已被修改'))
    return issues


def _selection_action(submitted: Optional[str], baseline_course: Optional[str]) -> str:
    if submitted is None or submitted == baseline_course:
        return 'no_change'
    if submitted == UNSELECTED_CODE:
        return 'set_unselected'
    return 'select'


def _parse_student_row(*, sheet: str, row: int, cells: Sequence[Optional[str]], roster: Mapping[str, WorkbookStudent], course_by_code: Mapping[str, WorkbookCourse], found: set[str], seen_global: dict[str, tuple[str, int]], issues: list[WorkbookIssue], year: str, term: str) -> Optional[WorkbookSelectionInput]:
    class_name, student_no, name, submitted = cells
    if not student_no:
        issues.append(WorkbookIssue(sheet, row, None, submitted, '学生行缺少学号'))
        return None
    student = roster.get(student_no)
    if student is None:
        issues.append(WorkbookIssue(sheet, row, student_no, submitted, '学生不属于该导出班级名册'))
        return None
    if student_no in found:
        issues.append(WorkbookIssue(sheet, row, student_no, submitted, '学生在同一班级页重复出现'))
        return None
    found.add(student_no)
    if student_no in seen_global:
        previous_sheet, previous_row = seen_global[student_no]
        issues.append(WorkbookIssue(sheet, row, student_no, submitted, f'学生跨工作表重复，首次出现在 {previous_sheet} 第{previous_row}行'))
        return None
    seen_global[student_no] = (sheet, row)
    if class_name != student.class_name or name != student.name:
        issues.append(WorkbookIssue(sheet, row, student_no, submitted, '学生班级或姓名与导出名册不一致'))
    if submitted and submitted != UNSELECTED_CODE and (submitted not in course_by_code):
        issues.append(WorkbookIssue(sheet, row, student_no, submitted, '课程代码不在导出目录中'))
    return WorkbookSelectionInput(sheet=sheet, row=row, year=year, term=term, grade=student.grade, class_name=student.class_name, student_no=student.student_no, name=student.name, baseline_course=student.baseline_course, submitted_course=submitted, action=_selection_action(submitted, student.baseline_course))


def _append_missing_roster_issues(sheet: str, roster: Mapping[str, WorkbookStudent], found: set[str], issues: list[WorkbookIssue]) -> None:
    for student_no in roster:
        if student_no not in found:
            issues.append(WorkbookIssue(sheet, None, student_no, None, '导出名册中的学生行缺失'))


def _parse_class_sheet(ws: Worksheet, sheet: str, roster: Mapping[str, WorkbookStudent], course_by_code: Mapping[str, WorkbookCourse], seen_global: dict[str, tuple[str, int]], issues: list[WorkbookIssue], year: str, term: str) -> list[WorkbookSelectionInput]:
    issues.extend(_validate_class_sheet_structure(ws, sheet))
    parsed_rows: list[WorkbookSelectionInput] = []
    found: set[str] = set()
    for row in range(2, ws.max_row + 1):
        cells = [normalize_cell(ws.cell(row=row, column=col).value) for col in range(1, 5)]
        if not any(cells):
            continue
        parsed = _parse_student_row(sheet=sheet, row=row, cells=cells, roster=roster, course_by_code=course_by_code, found=found, seen_global=seen_global, issues=issues, year=year, term=term)
        if parsed is not None:
            parsed_rows.append(parsed)
    _append_missing_roster_issues(sheet, roster, found, issues)
    return parsed_rows


def _parse_class_sheets(wb, expected_sheets: set[str], roster_by_sheet: Mapping[str, Mapping[str, WorkbookStudent]], course_by_code: Mapping[str, WorkbookCourse], issues: list[WorkbookIssue], year: str, term: str) -> tuple[list[WorkbookSelectionInput], dict[tuple[str, str], WorkbookSelectionInput]]:
    parsed_rows: list[WorkbookSelectionInput] = []
    seen_global: dict[str, tuple[str, int]] = {}
    for sheet in sorted(expected_sheets):
        if sheet not in wb.sheetnames:
            continue
        parsed_rows.extend(_parse_class_sheet(wb[sheet], sheet, roster_by_sheet[sheet], course_by_code, seen_global, issues, year, term))
    row_lookup = {(row.sheet, row.student_no): row for row in parsed_rows}
    return (parsed_rows, row_lookup)


def _final_course_code(student: WorkbookStudent, parsed: Optional[WorkbookSelectionInput], course_by_code: Mapping[str, WorkbookCourse]) -> Optional[str]:
    if parsed is None or parsed.submitted_course is None:
        final_code = student.baseline_course
    elif parsed.submitted_course == UNSELECTED_CODE:
        final_code = None
    elif parsed.submitted_course in course_by_code:
        final_code = parsed.submitted_course
    else:
        final_code = None
    return None if final_code == UNSELECTED_CODE else final_code


def _build_candidate_occupancy(students: Sequence[WorkbookStudent], row_lookup: Mapping[tuple[str, str], WorkbookSelectionInput], course_by_code: Mapping[str, WorkbookCourse]) -> Counter[tuple[str, str, str]]:
    occupancy: Counter[tuple[str, str, str]] = Counter()
    for student in students:
        parsed = row_lookup.get((student.sheet, student.student_no))
        final_code = _final_course_code(student, parsed, course_by_code)
        if final_code:
            occupancy[student.grade, student.class_name, final_code] += 1
    return occupancy


def _class_suggestions(students: Sequence[WorkbookStudent], courses: Sequence[WorkbookCourse], occupancy: Mapping[tuple[str, str, str], int]) -> dict[tuple[str, str], tuple[WorkbookSuggestion, ...]]:
    class_keys = {(student.grade, student.class_name) for student in students}
    return {key: _suggestions_for_class(key, courses, occupancy) for key in class_keys}


def _attach_issue_suggestions(issues: Sequence[WorkbookIssue], students: Sequence[WorkbookStudent], class_suggestions: Mapping[tuple[str, str], Sequence[WorkbookSuggestion]]) -> list[WorkbookIssue]:
    enhanced: list[WorkbookIssue] = []
    student_index = {(student.sheet, student.student_no): student for student in students}
    for issue in issues:
        student = student_index.get((issue.sheet, issue.student_no or ''))
        suggestions = class_suggestions.get((student.grade, student.class_name), ()) if student else ()
        enhanced.append(WorkbookIssue(issue.sheet, issue.row, issue.student_no, issue.course_code, issue.reason, issue.suggestions or suggestions))
    return enhanced


def _build_baseline_occupancy(students: Sequence[WorkbookStudent]) -> Counter[tuple[str, str, str]]:
    occupancy: Counter[tuple[str, str, str]] = Counter()
    for student in students:
        if student.baseline_course and student.baseline_course != UNSELECTED_CODE:
            occupancy[student.grade, student.class_name, student.baseline_course] += 1
    return occupancy


def _capacity_issue_rows(parsed_rows: Sequence[WorkbookSelectionInput], grade: str, class_name: str, code: str) -> Iterable[WorkbookSelectionInput]:
    return (parsed for parsed in parsed_rows if parsed.grade == grade and parsed.class_name == class_name and (parsed.submitted_course == code) and (parsed.baseline_course != code))


def _append_capacity_issues(issues: list[WorkbookIssue], parsed_rows: Sequence[WorkbookSelectionInput], course_by_code: Mapping[str, WorkbookCourse], occupancy: Mapping[tuple[str, str, str], int], baseline_occupancy: Mapping[tuple[str, str, str], int], class_suggestions: Mapping[tuple[str, str], Sequence[WorkbookSuggestion]]) -> None:
    for (grade, class_name, code), count in sorted(occupancy.items()):
        course = course_by_code.get(code)
        baseline_count = baseline_occupancy[grade, class_name, code]
        if course is None or count <= course.quota or count <= baseline_count:
            continue
        suggestions = class_suggestions.get((grade, class_name), ())
        for parsed in _capacity_issue_rows(parsed_rows, grade, class_name, code):
            issues.append(WorkbookIssue(parsed.sheet, parsed.row, parsed.student_no, code, f'本班候选最终人数 {count} 超过限额 {course.quota}；离线结果不预留名额', suggestions))


def _parse_open_workbook(wb, expected: Optional[WorkbookExpectation]) -> ParsedCourseSelectionWorkbook:
    year, term, baseline_id, students, courses, issues = _read_validated_metadata(wb)
    issues.extend(_validate_expected_contract(year, term, baseline_id, students, courses, expected))
    catalog = wb[CATALOG_SHEET] if CATALOG_SHEET in wb.sheetnames else None
    issues.extend(_validate_catalog(catalog, courses))
    expected_sheets = {student.sheet for student in students}
    issues.extend(_validate_workbook_sheet_set(wb, expected_sheets))
    roster_by_sheet = _build_roster_by_sheet(students, issues)
    course_by_code = {course.course_code: course for course in courses}
    parsed_rows, row_lookup = _parse_class_sheets(wb, expected_sheets, roster_by_sheet, course_by_code, issues, year, term)
    occupancy = _build_candidate_occupancy(students, row_lookup, course_by_code)
    suggestions = _class_suggestions(students, courses, occupancy)
    issues = _attach_issue_suggestions(issues, students, suggestions)
    _append_capacity_issues(issues, parsed_rows, course_by_code, occupancy, _build_baseline_occupancy(students), suggestions)
    if issues:
        raise WorkbookValidationError(issues)
    return ParsedCourseSelectionWorkbook(format=WORKBOOK_FORMAT, version=WORKBOOK_VERSION, year=year, term=term, baseline_id=baseline_id, rows=tuple(parsed_rows))


def parse_course_selection_workbook(content: bytes | BytesIO, *, expected: Optional[WorkbookExpectation]=None) -> ParsedCourseSelectionWorkbook:
    """Validate every actual class sheet and return a complete, DB-free input set.

    Blank cells mean ``no_change``. Explicit ``未选`` is a proposed cancellation.
    No conflict is auto-resolved and no capacity result here reserves a seat.
    """
    raw_content = _workbook_bytes(content)
    validate_xlsx_archive_budget(raw_content)
    wb = _load_course_selection_workbook(raw_content)
    try:
        return _parse_open_workbook(wb, expected)
    finally:
        wb.close()
